import sqlite3
import io
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient

from rollcall import courses, db, eligibility, evidence, register, roster, templates
from rollcall.app import create_app
from rollcall.device import RawPunch
from rollcall.punches import store_punches

L200 = ("matric_no,name,level\n"
        "A/1,ALPHA Ada,200\nA/2,BRAVO Bola,200\nA/3,CHARLIE Chidi,200\nA/4,DELTA Dayo,200\nA/5,ECHO Efe,200\n")
L300 = ("matric_no,name,level\n"
        "B/1,FOXTROT Femi,300\nB/2,GOLF Gbenga,300\nB/3,HOTEL Hauwa,300\nB/4,INDIA Ify,300\nB/5,JULIET Jide,300\n")


def level_id(conn, number):
    return conn.execute("SELECT id FROM levels WHERE level_number = ?", (number,)).fetchone()["id"]


def import_csv(conn, number, text):
    import_id = conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, 1, 1, ?, 'pending', ?)", (level_id(conn, number), text, db.now())).lastrowid
    roster.commit(conn, import_id, 1)
    # An import makes no course (2026-09-24); these tests use one named for the level.
    made = conn.execute("SELECT id FROM courses WHERE level_id = ? ORDER BY id", (level_id(conn, number),)).fetchone()
    return made["id"] if made else courses.create_course(conn, 1, level_id(conn, number), f"{number} Level", None, 1)


def person(conn, matric):
    return conn.execute("SELECT * FROM people WHERE matric_no = ?", (matric,)).fetchone()


def scan(conn, matrics, when):
    store_punches(conn, 1, [RawPunch(str(person(conn, m)["device_user_id"]), datetime.fromisoformat(when), 1, 255)
                            for m in matrics], "network", 1)


def status_of(reg, matric):
    return next(s for s in reg["students"] if s["matric_no"] == matric)["status"]


@pytest.fixture
def two_levels(conn):
    c200 = import_csv(conn, 200, L200)
    c300 = import_csv(conn, 300, L300)
    return c200, c300


# -- 1. courses and carry-over students ----------------------------------------------------------

def test_an_import_makes_no_course_and_a_course_made_later_takes_the_whole_level(conn):
    # Reported 2026-09-24: importing five levels made five courses called "200 Level" to "600 Level".
    # A course is the lecturer's to make, with its real code.
    lvl = level_id(conn, 200)
    import_id = conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, 1, 1, ?, 'pending', ?)", (lvl, L200, db.now())).lastrowid
    summary = roster.commit(conn, import_id, 1)
    assert conn.execute("SELECT COUNT(*) FROM courses").fetchone()[0] == 0
    assert summary["course_ids"] == [] and summary["template_id"] is None

    course = courses.create_course(conn, 1, lvl, "PCT 201", "Pharmaceutics", 1)
    assert [m["matric_no"] for m in courses.members(conn, course)] == ["A/1", "A/2", "A/3", "A/4", "A/5"]
    assert templates.get_template(conn, templates.ensure_template(conn, course))["course_code"] == "PCT 201"

    # A later list for the level puts its new names on the course, and still makes none of its own.
    import_id = conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, 1, 1, ?, 'pending', ?)", (lvl, L200 + "A/6,KILO Kemi,200\n", db.now())).lastrowid
    assert roster.commit(conn, import_id, 1)["course_ids"] == [course]
    assert conn.execute("SELECT COUNT(*) FROM courses").fetchone()[0] == 1
    assert "A/6" in [m["matric_no"] for m in courses.members(conn, course)]


def test_carry_over_preview_and_commit(conn, two_levels):
    c200, _ = two_levels
    entries = [
        {"matric_no": "B/1"},                                          # existing 300 level student
        {"matric_no": "C/9", "name": "KILO Kemi", "level": "400"},      # new to the app, 400 level
        {"matric_no": "A/1"},                                          # already on the course
        {"matric_no": "Z/1", "name": "LIMA Lola"},                      # new but no level given
        {"matric_no": "B/1"},                                          # duplicate in the list
    ]
    plan = courses.carry_over_plan(conn, c200, entries)
    assert (plan["enrol"], plan["create"], plan["already_on_course"], plan["skipped"]) == (1, 1, 1, 2)
    assert plan["entries"][1]["device_user_id"] == 4001 and plan["entries"][1]["kind"] == "carry_over"
    assert conn.execute("SELECT COUNT(*) FROM people WHERE matric_no = 'C/9'").fetchone()[0] == 0

    result = courses.carry_over_commit(conn, c200, entries, 1)
    assert result["create"] == 1
    members = courses.members(conn, c200)
    assert [m["matric_no"] for m in members] == ["A/1", "A/2", "A/3", "A/4", "A/5", "B/1", "C/9"]
    assert [m["kind"] for m in members][-2:] == ["carry_over", "carry_over"]
    assert (members[-1]["device_user_id"], members[-1]["level_number"]) == (4001, 400)


def test_new_carry_over_student_never_takes_an_id_seen_only_in_scans(conn, two_levels):
    from rollcall.punches import unmatched
    c200, _ = two_levels
    # 4001 is the first 400 Level ID, and a stranger has already scanned with it.
    store_punches(conn, 1, [RawPunch("4001", datetime(2026, 9, 14, 9, 5), 1, 255)], "network", 1)
    entries = [{"matric_no": "C/9", "name": "KILO Kemi", "level": "400"}]
    assert courses.carry_over_plan(conn, c200, entries)["entries"][0]["device_user_id"] == 4002
    courses.carry_over_commit(conn, c200, entries, 1)
    assert person(conn, "C/9")["device_user_id"] == 4002
    assert "4001" in [u["device_user_id"] for u in unmatched(conn)]


def test_new_level_students_join_every_course_of_their_level(conn, two_levels):
    c200, _ = two_levels
    c203 = courses.create_course(conn, 1, level_id(conn, 200), "PCT 203", "Physiology", 1)
    assert len(courses.members(conn, c203)) == 5
    import_csv(conn, 200, "matric_no,name,level\nA/6,MIKE Musa,200\n")
    assert "A/6" in [m["matric_no"] for m in courses.members(conn, c200)]
    assert "A/6" in [m["matric_no"] for m in courses.members(conn, c203)]


def test_removed_student_leaves_the_register_and_can_be_restored(conn, two_levels):
    c200, _ = two_levels
    s = register.create_session(conn, c200, "2026-09-15", "09:00", "11:00", "device", 1)
    with pytest.raises(courses.CourseError):
        courses.remove_student(conn, c200, person(conn, "A/5")["id"], "", 1)
    courses.remove_student(conn, c200, person(conn, "A/5")["id"], "Dropped the course", 1)
    assert "A/5" not in [x["matric_no"] for x in register.session_register(conn, s)["students"]]
    assert courses.carry_over_plan(conn, c200, [{"matric_no": "A/5"}])["entries"][0]["action"] == "restore"
    courses.carry_over_commit(conn, c200, [{"matric_no": "A/5"}], 1)
    assert "A/5" in [x["matric_no"] for x in register.session_register(conn, s)["students"]]


def test_carry_over_scans_count_and_do_not_make_a_second_class(conn, two_levels):
    c200, c300 = two_levels
    courses.carry_over_commit(conn, c200, [{"matric_no": m} for m in ("B/1", "B/2", "B/3", "B/4", "B/5")], 1)
    scan(conn, ["A/1", "A/2", "A/3", "A/4", "A/5"], "2026-09-15T10:50:00")
    scan(conn, ["B/1", "B/2", "B/3", "B/4", "B/5"], "2026-09-15T10:52:00")   # carry-over, same class
    created = register.detect_sessions(conn, 1)
    assert len(created) == 1 and created[0]["course_id"] == c200
    reg = register.session_register(conn, created[0]["id"])
    assert reg["totals"]["present"] == 10 and status_of(reg, "B/3") == "present"


def test_two_courses_with_the_same_students_leave_the_class_for_the_lecturer(conn, two_levels):
    c200, _ = two_levels
    c203 = courses.create_course(conn, 1, level_id(conn, 200), "PCT 203", None, 1)
    scan(conn, ["A/1", "A/2", "A/3", "A/4", "A/5"], "2026-09-16T12:00:00")
    created = register.detect_sessions(conn, 1)
    assert created[0]["needs_course"] is True
    unassigned = register.unassigned_sessions(conn)
    assert unassigned[0]["candidate_course_ids"] == [c200, c203]
    assert register.session_register(conn, created[0]["id"])["unassigned"] is True
    register.assign_course(conn, created[0]["id"], c203, 1)
    assert register.unassigned_sessions(conn) == []
    assert register.session_register(conn, created[0]["id"])["totals"]["present"] == 5
    assert register.detect_sessions(conn, 1) == []


# -- 2. excused absences -------------------------------------------------------------------------

def four_classes(conn, course):
    sessions = [register.create_session(conn, course, f"2026-09-{14 + d:02d}", "09:00", "11:00", "device", 1)
                for d in range(4)]
    for d in range(4):
        scan(conn, ["A/1"], f"2026-09-{14 + d:02d}T10:55:00")
    return sessions


def test_excused_absence_comes_out_of_the_total(conn, two_levels):
    c200, _ = two_levels
    s = four_classes(conn, c200)
    bola = person(conn, "A/2")["id"]
    scan(conn, ["A/2"], "2026-09-14T10:55:00")
    scan(conn, ["A/2"], "2026-09-17T10:55:00")
    excusal = register.add_excusal(conn, bola, "2026-09-15", "2026-09-16", "medical", "Clinic note 114", 1)

    assert status_of(register.session_register(conn, s[1]), "A/2") == "excused"
    summary = {r["matric_no"]: r for r in register.course_summary(conn, c200)["students"]}
    assert (summary["A/2"]["attended"], summary["A/2"]["excused"], summary["A/2"]["counted"],
            summary["A/2"]["percent"]) == (2, 2, 2, 100.0)

    # A scan inside an excused period still counts as present, and a correction beats both.
    scan(conn, ["A/2"], "2026-09-16T10:55:00")
    assert status_of(register.session_register(conn, s[2]), "A/2") == "present"
    register.add_mark(conn, s[2], bola, "absent", "Scanned and left", 1)
    assert status_of(register.session_register(conn, s[2]), "A/2") == "marked_absent"

    register.revoke_excusal(conn, excusal, "Clinic note not genuine", 1)
    assert status_of(register.session_register(conn, s[1]), "A/2") == "absent"
    assert register.list_excusals(conn, person_id=bola)[0]["revoke_note"] == "Clinic note not genuine"


def test_excusal_rules_and_scope(conn, two_levels):
    c200, _ = two_levels
    c203 = courses.create_course(conn, 1, level_id(conn, 200), "PCT 203", None, 1)
    s = four_classes(conn, c200)
    other = register.create_session(conn, c203, "2026-09-15", "13:00", "14:00", "device", 1)
    dayo = person(conn, "A/4")["id"]
    register.add_excusal(conn, dayo, "2026-09-14", "2026-09-17", "official", "Inter-faculty sports", 1, course_id=c200)
    assert status_of(register.session_register(conn, s[0]), "A/4") == "excused"
    assert status_of(register.session_register(conn, other), "A/4") == "absent"

    conn.execute("UPDATE settings SET value = 'present' WHERE key = 'excused_rule'")
    row = {r["matric_no"]: r for r in register.course_summary(conn, c200)["students"]}["A/4"]
    assert (row["counted"], row["percent"]) == (4, 100.0)

    for bad in (("2026-09-17", "2026-09-14", "medical", "x"), ("2026-09-14", "2026-09-14", "holiday", "x"),
                ("2026-09-14", "2026-09-14", "medical", " ")):
        with pytest.raises(register.RegisterError):
            register.add_excusal(conn, dayo, *bad, 1)


# -- 3. eligibility list -------------------------------------------------------------------------

def test_eligibility_decisions_and_rules(conn, two_levels):
    c200, _ = two_levels
    courses.carry_over_commit(conn, c200, [{"matric_no": "B/1"}], 1)
    four_classes(conn, c200)                               # A/1 attends all 4
    for d in (14, 15, 16):
        scan(conn, ["A/2", "B/1"], f"2026-09-{d}T10:55:00")   # 3 of 4 = 75%
    scan(conn, ["A/3"], "2026-09-14T10:55:00")              # 1 of 4 = 25%
    register.add_excusal(conn, person(conn, "A/5")["id"], "2026-09-14", "2026-09-17", "medical", "Admitted", 1)

    data = eligibility.eligibility(conn, c200)
    decisions = {s["matric_no"]: s["decision"] for s in data["students"]}
    assert decisions == {"A/1": "eligible", "A/2": "eligible", "A/3": "not_eligible", "A/4": "not_eligible",
                         "A/5": "review", "B/1": "eligible"}
    assert data["counts"] == {"total": 6, "eligible": 3, "not_eligible": 2, "review": 1, "carry_over": 1}
    assert data["rules"]["provisional"] is True                      # template plans 12 classes, 4 held
    assert data["rules"]["excused_text"].startswith("Excused absences are left out")
    assert next(s for s in data["students"] if s["matric_no"] == "B/1")["carry_over"] is True

    templates.set_weeks(conn, templates.ensure_template(conn, c200), 1, [{"classes": 1}] * 4)
    assert eligibility.eligibility(conn, c200)["rules"]["provisional"] is False


def test_eligibility_exports(conn, two_levels):
    from openpyxl import load_workbook

    c200, _ = two_levels
    four_classes(conn, c200)
    data = eligibility.eligibility(conn, c200)
    full = load_workbook(io.BytesIO(eligibility.to_xlsx(data, "full"))).active
    values = [c.value for row in full.iter_rows() for c in row if c.value is not None]
    assert "EXAMINATION ELIGIBILITY LIST (PROVISIONAL)" in values and "ALPHA Ada" in values
    assert any(isinstance(v, str) and v.startswith("HEAD OF DEPARTMENT") for v in values)

    notice = load_workbook(io.BytesIO(eligibility.to_xlsx(data, "notice"))).active
    notice_values = [str(c.value) for row in notice.iter_rows() for c in row if c.value is not None]
    assert "A/1" in notice_values and not any("ALPHA" in v for v in notice_values)

    for variant in ("full", "notice"):
        assert eligibility.to_pdf(data, variant)[:5] == b"%PDF-"
    assert eligibility.filename(data, "pdf", "notice") == "Eligibility 200 Level 2026-2027 notice.pdf"


# -- 4. student evidence -------------------------------------------------------------------------

def test_search_and_evidence(conn, two_levels):
    c200, _ = two_levels
    s = four_classes(conn, c200)
    bola = person(conn, "A/2")
    scan(conn, ["A/2"], "2026-09-14T10:41:07")
    register.add_mark(conn, s[1], bola["id"], "present", "Signed paper register, device queue", 1)
    register.add_excusal(conn, bola["id"], "2026-09-16", "2026-09-16", "bereavement", "Family funeral", 1)

    assert [r["matric_no"] for r in evidence.search(conn, "A/2")] == ["A/2"]
    assert [r["matric_no"] for r in evidence.search(conn, "bravo")] == ["A/2"]
    # Accents and case beyond ASCII: SQLite's LIKE alone found neither of these.
    import_csv(conn, 200, "matric_no,name,level\nY/9,\u1eccL\u00c1B\u00cdS\u00cd Oluwa\u1e63eun,200\n")
    assert [r["matric_no"] for r in evidence.search(conn, "olabisi")] == ["Y/9"]
    assert [r["matric_no"] for r in evidence.search(conn, "\u1ecdl\u00e1b\u00ed")] == ["Y/9"]     # "ọlábí"
    assert [r["matric_no"] for r in evidence.search(conn, "oluwaseun")] == ["Y/9"]
    assert evidence.search(conn, str(bola["device_user_id"]))[0]["matric_no"] == "A/2"
    with pytest.raises(evidence.EvidenceError):
        evidence.search(conn, "a")

    ev = evidence.student_evidence(conn, bola["id"])
    assert ev["student"]["matric_no"] == "A/2" and ev["levels"][0]["level"] == "200 Level"
    assert ev["scans"][0]["timestamp"] == "2026-09-14T10:41:07" and ev["scans"][0]["source"] == "network"
    course = ev["courses"][0]
    reasons = [c["reason"] for c in course["classes"]]
    assert reasons[0].startswith("Scanned at 10:41:07")
    assert reasons[1].startswith("Corrected to present by")
    assert reasons[2] == "Excused (bereavement): Family funeral"
    assert reasons[3].startswith("No scan between 08:45 and 11:30")
    assert course["corrections"][0]["note"] == "Signed paper register, device queue"
    assert course["summary"]["attended"] == 2 and course["summary"]["excused"] == 1


def test_api_for_courses_excusals_eligibility_and_evidence(tmp_path, k40_factory, fake):
    with TestClient(create_app(tmp_path / "e.db", k40_factory=k40_factory)) as c:
        c.post("/api/setup", json={"username": "lecturer", "display_name": "Dr", "password": "secret123",
                                   "academic_session": "2026/2027"})
        levels = {l["level_number"]: l["id"] for l in c.get("/api/levels").json()}
        preview = c.post(f"/api/levels/{levels[200]}/roster/preview", files={"file": ("r.csv", L200.encode())}).json()
        c.post(f"/api/roster-imports/{preview['import_id']}/commit")
        c.post("/api/courses", json={"level_id": levels[200], "code": "200 Level"})
        course = c.get("/api/courses").json()[0]
        assert c.patch(f"/api/courses/{course['id']}", json={"code": "PCT 201", "title": "Pharmacology"}).json()["code"] == "PCT 201"

        csv_text = "matric_no,name,level\nX/1,NOVEMBER Ngozi,300\nX/2,OSCAR Obi,\n"
        plan = c.post(f"/api/courses/{course['id']}/students/preview", files={"file": ("co.csv", csv_text.encode())}).json()
        assert (plan["create"], plan["skipped"]) == (1, 1)
        json_plan = c.post(f"/api/courses/{course['id']}/students/preview",
                           json={"entries": [{"matric_no": "X/1", "name": "NOVEMBER Ngozi", "level": 300}]}).json()
        assert json_plan["create"] == 1
        added = c.post(f"/api/courses/{course['id']}/students",
                       json={"entries": [{"matric_no": "X/1", "name": "NOVEMBER Ngozi", "level": 300}]}).json()
        assert added["added"][0]["device_user_id"] == 3001
        students = c.get(f"/api/courses/{course['id']}/students").json()
        assert students[-1]["kind"] == "carry_over" and students[-1]["level"] == 300

        push = c.post("/api/devices/1/push", json={"course_id": course["id"]}).json()
        assert push["pushed"] == 6 and 3001 in fake.users

        session = c.post("/api/sessions", json={"course_id": course["id"], "date": "2026-09-15",
                                                 "start_time": "09:00", "end_time": "11:00"}).json()["id"]
        ngozi = students[-1]["person_id"]
        ex = c.post("/api/excusals", json={"person_id": ngozi, "date_from": "2026-09-15", "date_to": "2026-09-15",
                                           "category": "medical", "note": "Hospital", "course_id": course["id"]})
        assert ex.status_code == 200
        assert c.post("/api/excusals", json={"person_id": ngozi, "date_from": "2026-09-15", "date_to": "2026-09-15",
                                             "category": "holiday", "note": "x"}).status_code == 422
        reg = c.get(f"/api/sessions/{session}/register").json()
        assert reg["totals"]["excused"] == 1

        elig = c.get(f"/api/courses/{course['id']}/eligibility").json()
        assert elig["counts"]["review"] == 1 and elig["course"]["code"] == "PCT 201"
        pdf = c.get(f"/api/courses/{course['id']}/eligibility/export", params={"format": "pdf", "variant": "notice"})
        assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
        assert "Eligibility PCT 201 2026-2027 notice.pdf" in pdf.headers["content-disposition"]

        found = c.get("/api/students/search", params={"q": "ngozi"}).json()
        ev = c.get(f"/api/students/{found[0]['id']}/evidence").json()
        assert ev["courses"][0]["classes"][0]["status"] == "excused"
        assert c.get("/api/sessions/unassigned").json() == []
        assert c.put("/api/settings", json={"excused_rule": "sometimes"}).status_code == 400
        # A setting added after the database was made has no row, as in a database from an older version:
        # it is listed with its default, and it saves.
        with sqlite3.connect(tmp_path / "e.db") as raw:
            raw.execute("DELETE FROM settings WHERE key = 'paper_size'")
        assert c.get("/api/settings").json()["paper_size"] == "a4"
        conn_rows = c.put("/api/settings", json={"paper_size": "letter"})
        assert conn_rows.status_code == 200 and conn_rows.json()["paper_size"] == "letter"
        assert c.get("/api/settings").json()["paper_size"] == "letter"
        assert c.put("/api/settings", json={"paper_size": "a3"}).status_code == 400


# -- migration -----------------------------------------------------------------------------------

def test_version_3_data_moves_onto_courses(tmp_path):
    conn = db.connect(tmp_path / "v3.db")
    db.initialise(conn, target=3)
    stamp = db.now()
    conn.execute("INSERT INTO academic_sessions(name, is_current, created_at) VALUES ('2026/2027', 1, ?)", (stamp,))
    conn.execute("INSERT INTO devices(name, ip, is_default, created_at) VALUES ('K40', '1.1.1.1', 1, ?)", (stamp,))
    conn.execute("INSERT INTO accounts(username, display_name, role, password_hash, created_at) VALUES ('l','L','admin','x',?)", (stamp,))
    lvl = level_id(conn, 200)
    pid = conn.execute("INSERT INTO people(role, device_user_id, device_id, matric_no, surname, device_name, created_at)"
                       " VALUES ('student', 2001, 1, 'A/1', 'ALPHA', 'ALPHA', ?)", (stamp,)).lastrowid
    conn.execute("INSERT INTO level_memberships VALUES (?, 1, ?)", (pid, lvl))
    tid = conn.execute("INSERT INTO register_templates(level_id, academic_session_id, title, course_code, week1_start,"
                       " created_at, updated_at) VALUES (?, 1, 'T', 'PCT 201', '2026-09-14', ?, ?)", (lvl, stamp, stamp)).lastrowid
    conn.execute("INSERT INTO register_weeks VALUES (?, 1, '2026-09-14', 2, NULL)", (tid,))
    sid = conn.execute("INSERT INTO class_sessions(level_id, academic_session_id, date, start_time, end_time, kind, created_at)"
                       " VALUES (?, 1, '2026-09-15', '09:00', '11:00', 'device', ?)", (lvl, stamp)).lastrowid
    conn.execute("INSERT INTO attendance_plans(level_id, academic_session_id, threshold_percent, updated_at) VALUES (?, 1, 70, ?)",
                 (lvl, stamp))

    db.initialise(conn)
    course = conn.execute("SELECT * FROM courses").fetchone()
    assert course["code"] == "PCT 201"
    assert conn.execute("SELECT course_id FROM class_sessions WHERE id = ?", (sid,)).fetchone()[0] == course["id"]
    assert conn.execute("SELECT course_id FROM register_templates WHERE id = ?", (tid,)).fetchone()[0] == course["id"]
    assert conn.execute("SELECT threshold_percent FROM attendance_plans WHERE course_id = ?", (course["id"],)).fetchone()[0] == 70
    assert [m["matric_no"] for m in courses.members(conn, course["id"])] == ["A/1"]
    assert templates.get_template(conn, tid)["total_classes"] == 2
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
    assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    conn.close()
