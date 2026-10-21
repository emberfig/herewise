import io
import sqlite3
from datetime import datetime

import pytest

from conftest import ROOT, ROSTER
from rollcall import courses, db, levels, roster
from rollcall.device import RawPunch
from rollcall.punches import store_punches


def level(conn, number):
    return conn.execute("SELECT * FROM levels WHERE level_number = ?", (number,)).fetchone()


def pending(conn, level_id, text, academic_id=1):
    return conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, ?, 1, ?, 'pending', ?)", (level_id, academic_id, text, db.now())).lastrowid


def test_real_roster_preview(conn):
    text = ROSTER.read_text(encoding="utf-8-sig")
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv(text))
    assert plan["new"] == 258
    assert plan["skipped"] == 0
    assert plan["id_range"] == [2001, 2258]
    assert plan["capacity"]["after_import"] == 258
    assert plan["can_commit"]
    assert any(s["after"] == "EFIKOSLI Opazsitofo E." for s in plan["shortened"])
    # Preview changes nothing.
    assert conn.execute("SELECT COUNT(*) FROM people").fetchone()[0] == 0


def test_commit_then_reimport_is_unchanged(conn):
    text = ROSTER.read_text(encoding="utf-8-sig")
    lvl = level(conn, 200)["id"]
    summary = roster.commit(conn, pending(conn, lvl, text), 1)
    assert summary["new"] == 258
    again = roster.build_plan(conn, lvl, 1, 1, roster.parse_csv(text))
    assert again["new"] == 0 and again["unchanged"] == 258 and not again["can_commit"]


def test_next_session_moves_level_and_keeps_device_id(conn):
    text = "matric_no,name,level\n23/PHARX01/015,ASUKAGKI Joshua Adebayo,200\n"
    roster.commit(conn, pending(conn, level(conn, 200)["id"], text), 1)
    person = conn.execute("SELECT * FROM people WHERE matric_no = '23/PHARX01/015'").fetchone()
    assert person["device_user_id"] == 2001

    conn.execute("UPDATE academic_sessions SET is_current = 0")
    new_session = conn.execute("INSERT INTO academic_sessions(name, is_current, created_at) VALUES ('2027/2028', 1, ?)",
                               (db.now(),)).lastrowid
    moved = text.replace(",200", ",300")
    plan = roster.build_plan(conn, level(conn, 300)["id"], new_session, 1, roster.parse_csv(moved))
    assert plan["students"][0]["action"] == "update"
    assert plan["students"][0]["device_user_id"] == 2001
    roster.commit(conn, pending(conn, level(conn, 300)["id"], moved, new_session), 1)
    assert conn.execute("SELECT device_user_id FROM people WHERE id = ?", (person["id"],)).fetchone()[0] == 2001
    assert conn.execute("SELECT COUNT(*) FROM level_memberships WHERE person_id = ?", (person["id"],)).fetchone()[0] == 2


def test_duplicates_and_wrong_level_rows_are_skipped(conn):
    text = ("matric_no,name,level\n"
            "A/1,ONE Person,200\n"
            "A/1,ONE Again,200\n"
            "A/2,TWO Person,300\n"
            ",NO Matric,200\n")
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv(text))
    # The row with no matric number is imported since 2026-09-23: direct-entry students have none
    # yet, and it is added later. Only the duplicate and the wrong level are skipped.
    assert plan["new"] == 2
    assert plan["skipped"] == 2
    no_matric = next(s for s in plan["students"] if s["surname"] == "NO")
    assert no_matric["action"] == "new" and any("direct entry (DE)" in i for i in no_matric["issues"])


def test_unquoted_comma_in_name_is_rejoined(conn):
    text = ("matric_no,name,level\n"
            "A/1,OZAN-DELE Emmanuel Chibueze, Xuziil,200\n"
            "A/2,GOSAFU` Wavasloguifoze Cbsao,200\n")
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv(text))
    assert plan["new"] == 2
    assert plan["students"][0]["other_names"] == "Emmanuel Chibueze Xuziil"
    assert plan["students"][1]["surname"] == "GOSAFU"


def test_capacity_blocks_commit(conn):
    conn.execute("UPDATE devices SET user_capacity = 2")
    text = "matric_no,name,level\nA/1,ONE P,200\nA/2,TWO P,200\nA/3,THREE P,200\n"
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv(text))
    assert not plan["can_commit"]
    assert "holds 2" in plan["errors"][0]


def test_ids_are_never_reused(conn):
    roster.commit(conn, pending(conn, level(conn, 200)["id"], "matric_no,name\nA/1,ONE P\nA/2,TWO P\n"), 1)
    conn.execute("UPDATE people SET active = 0 WHERE matric_no = 'A/2'")
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv("matric_no,name\nA/3,THREE P\n"))
    assert plan["students"][0]["device_user_id"] == 2003


def unmatched_ids(conn):
    from rollcall.punches import unmatched
    return [u["device_user_id"] for u in unmatched(conn)]


def test_an_id_seen_only_in_scans_is_never_given_to_a_new_student(conn):
    roster.commit(conn, pending(conn, level(conn, 200)["id"], "matric_no,name\nA/1,ONE P\n"), 1)
    # A stranger on the terminal: ID 2002 scanned, but nobody here has it. Text IDs that are not whole
    # numbers, or sit in another block, must not move the count.
    store_punches(conn, 1, [RawPunch("2002", datetime(2026, 9, 14, 9, 5), 1, 255),
                            RawPunch("20x9", datetime(2026, 9, 14, 9, 6), 1, 255),
                            RawPunch("3150", datetime(2026, 9, 14, 9, 7), 1, 255)], "network", 1)
    assert "2002" in unmatched_ids(conn)

    import_id = pending(conn, level(conn, 200)["id"], "matric_no,name\nA/2,TWO P\n")
    plan = roster.build_plan(conn, level(conn, 200)["id"], 1, 1, roster.parse_csv("matric_no,name\nA/2,TWO P\n"))
    assert plan["id_range"] == [2003, 2003]
    roster.commit(conn, import_id, 1)
    assert conn.execute("SELECT device_user_id FROM people WHERE matric_no = 'A/2'").fetchone()[0] == 2003
    # The stranger's scan still belongs to nobody.
    assert "2002" in unmatched_ids(conn)


def test_punches_are_idempotent_and_immutable(conn):
    records = [RawPunch("2001", datetime(2026, 9, 15, 9, 14, 23), 1, 255),
               RawPunch("7777", datetime(2026, 9, 15, 9, 20, 0), 1, 255)]
    first = store_punches(conn, 1, records, "network", 1)
    second = store_punches(conn, 1, records, "usb", 1, file_name="attlog.dat")
    assert (first["new"], second["new"], second["already_known"]) == (2, 0, 2)
    assert conn.execute("SELECT COUNT(*) FROM punches").fetchone()[0] == 2
    with pytest.raises(sqlite3.DatabaseError, match="immutable"):
        conn.execute("UPDATE punches SET timestamp = '2026-09-15T10:00:00'")
    with pytest.raises(sqlite3.DatabaseError, match="immutable"):
        conn.execute("DELETE FROM punches")


def test_audit_log_is_append_only(conn):
    db.audit(conn, 1, "test")
    with pytest.raises(sqlite3.DatabaseError, match="append-only"):
        conn.execute("DELETE FROM audit_log")


# -- the file the portal actually hands out ------------------------------------------------------
#
# A first run against a real course list from the university on 2026-09-21 could not import it, for two reasons
# that had nothing to do with the roster itself: the file is a binary .xls, which was refused
# outright, and its second column is headed "Matric No." with a full stop, which did not match the
# column name the parser required even once the file was saved as CSV.

PORTAL_XLS = ROOT / "data" / "portal_roster.xls"


def portal_rows():
    """The shape the portal exports: a serial column, "Matric No." with a full stop, and the
    sheet's unused columns trailing behind every row."""
    return [
        ["S/N", "Matric No.", "Name", "", "", ""],
        ["1", "99/TEST01/001", "ADEMOLA Bunmi Oluwatoyin", "", "", ""],
        ["2", "99/TEST01/002", "CHUKWU Ngozi Adaeze", "", "", ""],
    ]


def test_the_column_the_portal_calls_matric_no_is_understood():
    rows = roster.parse_rows(portal_rows())
    assert [r["matric_no"] for r in rows] == ["99/TEST01/001", "99/TEST01/002"]
    assert rows[0]["name"] == "ADEMOLA Bunmi Oluwatoyin"


def test_a_sheets_unused_columns_are_not_folded_into_a_name():
    """Rows longer than the header are treated as a name with an unquoted comma in it, which is
    right for "OZAN-DELE Emmanuel, Xuziil" and wrong for the five empty columns a spreadsheet
    carries with it."""
    rows = roster.parse_rows([["Matric No.", "Name"],
                              ["Z/1", "SANI Halima", "", "", "", ""]])
    assert rows[0]["name"] == "SANI Halima"
    assert "," not in rows[0]["name"]


def test_every_spelling_of_the_matric_column_is_understood():
    for header in ("Matric No.", "MATRIC NO", "Matric  Number", "matric_no", "Matriculation No.",
                   "Registration Number"):
        rows = roster.parse_rows([[header, "Full Name"], ["Z/1", "SANI Halima"]])
        assert rows[0]["matric_no"] == "Z/1", header
        assert rows[0]["name"] == "SANI Halima", header


def test_a_header_that_names_nothing_useful_says_what_it_found():
    with pytest.raises(roster.RosterError) as raised:
        roster.parse_rows([["S/N", "Student ID", "Surname"], ["1", "Z/1", "SANI"]])
    said = str(raised.value)
    assert "matric number column" in said and "S/N, Student ID, Surname" in said


def test_a_roster_saved_as_the_old_xls_imports():
    rows = roster.parse_upload(PORTAL_XLS.read_bytes(), "PCT 999-Full-Time-course-list (1).xls")
    assert len(rows) == 5
    assert rows[0]["matric_no"] == "99/TEST01/001"
    assert rows[3]["name"] == "OYELARAN-ADEBOWALE Oluwatimilehin Adeola"
    # xlrd hands back every number as a float, and "200.0" is not a level.
    assert rows[0]["level"] == "200"


def test_a_roster_saved_as_xlsx_imports():
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    for row in (["S/N", "Matric No.", "Name", "Level"],
                [1, "99/TEST01/007", "OKAFOR Chinaza", 200.0]):
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)

    rows = roster.parse_upload(out.getvalue(), "roster.xlsx")
    assert rows[0]["matric_no"] == "99/TEST01/007"
    assert rows[0]["level"] == "200"


def test_a_web_page_named_xls_says_what_it_really_is():
    page = b"<html><head><meta charset=utf-8></head><body><table><tr><td>Matric</td></tr></table>"
    with pytest.raises(roster.RosterError) as raised:
        roster.parse_upload(page, "course-list.xls")
    said = str(raised.value)
    assert "web page" in said and "course-list.xls" in said


def test_a_spreadsheet_is_kept_with_its_import_as_the_rows_that_were_read(conn):
    """The commit re-reads the text stored with the import, so a spreadsheet has to be stored as
    the rows the preview showed. Stored as raw bytes it would commit nothing at all."""
    rows = roster.parse_upload(PORTAL_XLS.read_bytes(), "portal_roster.xls")
    text = roster.as_csv(rows)
    assert text.splitlines()[0] == "matric_no,name,level,surname,other_names"

    summary = roster.commit(conn, pending(conn, level(conn, 200)["id"], text), 1)
    assert summary["new"] == 5
    saved = conn.execute(
        "SELECT matric_no, surname FROM people WHERE matric_no LIKE '99/TEST01/%' ORDER BY matric_no"
    ).fetchall()
    assert [r["matric_no"] for r in saved] == [f"99/TEST01/00{n}" for n in range(1, 6)]
    assert saved[3]["surname"] == "OYELARAN-ADEBOWALE"


# -- the years a department actually runs --------------------------------------------------------
#
# Seeded 100 to 500 with no way to add or remove one, so a department stopping at 400 and a course
# of study running to 700 could neither be set up nor tidied up (found on a first run, 2026-09-21).

def test_a_department_that_runs_to_700_can_add_the_years(conn):
    added = levels.add(conn, 600, 1)
    assert added["name"] == "600 Level"
    assert (added["id_block_start"], added["id_block_end"]) == (6001, 6999)
    seventh = levels.add(conn, 700, 1, name="Final Year")
    assert seventh["name"] == "Final Year"
    assert seventh["level_number"] == 700
    assert (seventh["id_block_start"], seventh["id_block_end"]) == (7001, 7999)

    numbers = [r["level_number"] for r in
               conn.execute("SELECT level_number FROM levels ORDER BY level_number")]
    assert numbers == [100, 200, 300, 400, 500, 600, 700]


def test_a_level_that_would_share_device_ids_is_refused(conn):
    with pytest.raises(levels.LevelError, match="already a 300 Level"):
        levels.add(conn, 300, 1)
    for bad, said in ((250, "goes up in hundreds"), (0, "runs from 100"), (1000, "runs from 100")):
        with pytest.raises(levels.LevelError, match=said):
            levels.add(conn, bad, 1)


def test_a_department_that_stops_at_400_can_remove_the_rest(conn):
    empty = conn.execute("SELECT id FROM levels WHERE level_number = 500").fetchone()["id"]
    assert levels.remove(conn, empty, 1) == {"removed": "500 Level"}
    assert not conn.execute("SELECT 1 FROM levels WHERE level_number = 500").fetchone()


def test_a_level_with_anything_on_it_is_never_removed(conn):
    text = "matric_no,name\nZ/1,SANI Halima\nZ/2,OKAFOR Chinaza\n"
    lvl = level(conn, 200)["id"]
    roster.commit(conn, pending(conn, lvl, text), 1)
    courses.create_course(conn, 1, lvl, "PCT 201", None, 1)

    with pytest.raises(levels.LevelError) as raised:
        levels.remove(conn, lvl, 1)
    said = str(raised.value)
    # It says what is on it, counted and in words, rather than only refusing.
    assert "cannot be removed" in said
    # Plain nouns in the list, with the nuance in its own sentence after it: the list is also
    # read inside a table cell, where a clause per item made the row 220px tall (2026-09-22).
    assert "2 students" in said
    assert "counted for this session and any past one" in said
    assert "1 course" in said and "1 roster import" in said
    assert conn.execute("SELECT 1 FROM levels WHERE id = ?", (lvl,)).fetchone()


def test_what_points_at_a_level_is_read_from_the_schema_itself(conn):
    """A table added later that references levels is covered without anybody remembering to add it
    to a list here."""
    lvl = level(conn, 200)["id"]
    assert levels.references(conn, lvl) == {}
    conn.execute("CREATE TABLE a_later_table(id INTEGER PRIMARY KEY,"
                 " level_id INTEGER NOT NULL REFERENCES levels(id))")
    conn.execute("INSERT INTO a_later_table(level_id) VALUES (?)", (lvl,))

    assert levels.references(conn, lvl) == {"a_later_table": 1}
    with pytest.raises(levels.LevelError, match="cannot be removed"):
        levels.remove(conn, lvl, 1)


def test_adding_and_removing_a_level_is_in_the_activity_log(conn):
    from rollcall import audit

    added = levels.add(conn, 600, 1)
    levels.remove(conn, added["id"], 1)
    said = [audit.sentence(r["action"], r["detail"] or "")
            for r in conn.execute("SELECT action, detail FROM audit_log WHERE action LIKE 'level.%'"
                                  " ORDER BY id")]
    assert said[0] == "Added a level: 600 Level, device IDs 6001 to 6999"
    assert said[1] == "Removed a level: 600 Level, which had nothing on it"


# -- taking an import back out -------------------------------------------------------------------

def test_an_import_into_the_wrong_level_can_be_taken_back(conn):
    """Asked for on a first run: a roster went into the wrong level and there was no way back short
    of editing the database (2026-09-21)."""
    text = "matric_no,name\nW/1,SANI Halima\nW/2,OKAFOR Chinaza\n"
    lvl = level(conn, 400)["id"]
    import_id = pending(conn, lvl, text)
    roster.commit(conn, import_id, 1)
    assert conn.execute("SELECT COUNT(*) FROM people WHERE roster_import_id = ?",
                        (import_id,)).fetchone()[0] == 2

    listed = roster.imports_for(conn, lvl)
    assert listed[0]["people_created"] == 2 and listed[0]["can_undo"] is True

    done = roster.undo(conn, import_id, 1)

    assert done["removed"] == 2
    assert not conn.execute("SELECT 1 FROM people WHERE matric_no LIKE 'W/%'").fetchone()
    assert conn.execute("SELECT status FROM roster_imports WHERE id = ?",
                        (import_id,)).fetchone()["status"] == "discarded"
    with pytest.raises(roster.RosterError, match="discarded"):
        roster.undo(conn, import_id, 1)


def test_an_import_whose_students_have_scanned_is_never_taken_back(conn):
    text = "matric_no,name\nW/1,SANI Halima\n"
    lvl = level(conn, 400)["id"]
    import_id = pending(conn, lvl, text)
    roster.commit(conn, import_id, 1)
    person = conn.execute("SELECT id, device_user_id FROM people WHERE matric_no = 'W/1'").fetchone()
    store_punches(conn, 1, [RawPunch(str(person["device_user_id"]), datetime(2026, 9, 15, 9, 0, 0), 1, 255)],
                  "network", 1)

    listed = roster.imports_for(conn, lvl)
    assert listed[0]["can_undo"] is False
    assert listed[0]["holding"]["scans"] == 1

    with pytest.raises(roster.RosterError) as raised:
        roster.undo(conn, import_id, 1)
    said = str(raised.value)
    assert "1 scans" in said and "never given to anybody else" in said
    assert conn.execute("SELECT 1 FROM people WHERE matric_no = 'W/1'").fetchone()


# -- a department's own workbook -----------------------------------------------------------------
#
# The portal hands out one list with the header on the first row. A department keeps every year in
# one workbook: a sheet per level, each opening with the university's name, the session and a
# title, and the columns named five or six rows down. Herewise refused all five sheets of the real
# one, on the first row it read (reported 2026-09-22). The names below are invented; the shape is
# the shape of that file.

def a_workbook(sheets: dict[str, list[list]]) -> bytes:
    import io

    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def fixture_surname(level: int, n: int) -> str:
    """SURNAMECAAB for 200 Level's first student: the level and number in letters, because a name with
    a digit in it is refused as not a name (2026-09-24)."""
    return "SURNAME" + "".join(chr(ord("A") + int(d)) for d in f"{level}{n}")


def a_level_sheet(level: int, students: int, start: int = 1) -> list[list]:
    """Title block, a blank row, the header, then the students: the real file's shape."""
    return [
        ["EXAMPLE UNIVERSITY, EXAMPLE TOWN, COLLEGE OF PHARMACY,"],
        ["  2026/2027 ACADEMIC SESSION"],
        ["", f"{level} LEVEL STUDENTS' PRACTICAL GROUPINGS"],
        [],
        [],
        ["S/N", "NAMES", "MATRIC NO.", "GROUP"],
        *[[str(n), f"{fixture_surname(level, n)} Given Other", f"25/PHARM01/{level}{n:03d}", "A"]
          for n in range(start, start + students)],
    ]


def test_a_header_below_a_title_block_is_found(conn):
    raw = a_workbook({"200 Level": a_level_sheet(200, 3)})
    rows = roster.parse_upload(raw, "Level groupings.xlsx", 200)
    assert len(rows) == 3
    assert rows[0]["matric_no"] == "25/PHARM01/200001"
    assert rows[0]["name"].startswith(fixture_surname(200, 1))


def test_the_line_numbers_are_the_ones_in_the_file(conn):
    """Somebody reading "line 8" has to find it at row 8 of their own spreadsheet, not eight rows
    after a header they cannot see."""
    raw = a_workbook({"200 Level": a_level_sheet(200, 3)})
    rows = roster.parse_upload(raw, "Level groupings.xlsx", 200)
    assert [r["_line"] for r in rows] == [7, 8, 9]


def test_the_sheet_for_the_level_being_imported_is_the_one_read(conn):
    """One workbook, a sheet per year. Taking the first sheet would have put 200 Level's students
    into whichever level the person was on."""
    raw = a_workbook({
        "200 Level": a_level_sheet(200, 2),
        "300L": a_level_sheet(300, 4),
        "400 Level ": a_level_sheet(400, 3),
    })
    assert len(roster.parse_upload(raw, "groupings.xlsx", 200)) == 2
    assert len(roster.parse_upload(raw, "groupings.xlsx", 300)) == 4, "300L is named for 300"
    assert len(roster.parse_upload(raw, "groupings.xlsx", 400)) == 3, "a trailing space is nothing"
    got = roster.parse_upload(raw, "groupings.xlsx", 300)
    assert all(r["matric_no"].startswith("25/PHARM01/300") for r in got)


def test_a_workbook_with_no_sheet_for_this_level_says_which_it_has(conn):
    raw = a_workbook({"200 Level": a_level_sheet(200, 2), "300L": a_level_sheet(300, 2)})
    with pytest.raises(roster.RosterError) as raised:
        roster.parse_upload(raw, "groupings.xlsx", 100)
    said = str(raised.value)
    assert "none of them is named for 100 Level" in said
    assert "200 Level" in said and "300L" in said


def test_one_sheet_is_read_whatever_it_is_called(conn):
    """A single-sheet file is the portal's, and its name means nothing."""
    raw = a_workbook({"Sheet1": a_level_sheet(200, 2)})
    assert len(roster.parse_upload(raw, "portal.xlsx", 500)) == 2


def test_a_file_that_names_no_columns_anywhere_is_still_refused(conn):
    raw = a_workbook({"Sheet1": [["EXAMPLE UNIVERSITY"], ["2026/2027"], ["1", "somebody"]]})
    with pytest.raises(roster.RosterError, match="names the matric number or the name column"):
        roster.parse_upload(raw, "nothing.xlsx", 200)


# -- direct-entry students, who have no matric number yet ----------------------------------------
#
# Reported 2026-09-23: the rows with no matric number in the department's list are direct-entry
# students who have not been given one yet. They are imported, recognised again by name within the
# level, and given their number later without becoming a second student.


def people(conn):
    return {f"{r['surname']} {r['other_names']}".strip(): dict(r)
            for r in conn.execute("SELECT * FROM people WHERE role = 'student'")}


def test_a_direct_entry_student_keeps_their_device_id_when_their_matric_number_comes(conn):
    lvl = level(conn, 200)["id"]
    roster.commit(conn, pending(conn, lvl, "matric_no,name\nA/1,ONE Person\n,ADEYEMI Direct Entry\n"), 1)
    before = people(conn)["ADEYEMI Direct Entry"]
    assert before["matric_no"] is None, "stored as an empty string, which a second one would collide with"

    # The same list again: the student with no number is recognised by name, not made twice.
    again = roster.build_plan(conn, lvl, 1, 1, roster.parse_csv(
        "matric_no,name\nA/1,ONE Person\n,ADEYEMI Direct Entry\n"))
    assert again["new"] == 0 and again["unchanged"] == 2

    # A later list carries the number: it is filled in on the same student.
    roster.commit(conn, pending(conn, lvl, "matric_no,name\nA/1,ONE Person\nA/9,ADEYEMI Direct Entry\n"), 1)
    after = people(conn)["ADEYEMI Direct Entry"]
    assert after["id"] == before["id"]
    assert after["device_user_id"] == before["device_user_id"], "a device ID is permanent"
    assert after["matric_no"] == "A/9"
    assert conn.execute("SELECT COUNT(*) FROM people").fetchone()[0] == 2


def test_two_students_without_matric_numbers_do_not_collide(conn):
    lvl = level(conn, 200)["id"]
    roster.commit(conn, pending(conn, lvl, "matric_no,name\n,FIRST Direct\n,SECOND Direct\n"), 1)
    got = people(conn)
    assert got["FIRST Direct"]["matric_no"] is None and got["SECOND Direct"]["matric_no"] is None


def test_a_name_or_matric_number_can_be_corrected_and_the_number_added(conn):
    lvl = level(conn, 200)["id"]
    roster.commit(conn, pending(conn, lvl, "matric_no,name\nA/1,ONE Person\n,ADEYEMI Direct Entry\n"), 1)
    student = people(conn)["ADEYEMI Direct Entry"]

    done = roster.edit_student(conn, student["id"], "ADEYEMI", "Direct Entry", "25/PHARM01/120", 1)
    assert done["changed"] and not done["device_name_changed"]
    assert people(conn)["ADEYEMI Direct Entry"]["matric_no"] == "25/PHARM01/120"

    renamed = roster.edit_student(conn, student["id"], "ADEYEMI-COLE", "Direct Entry", "25/PHARM01/120", 1)
    assert renamed["device_name_changed"]
    row = conn.execute("SELECT * FROM people WHERE id = ?", (student["id"],)).fetchone()
    assert row["surname"] == "ADEYEMI-COLE" and row["device_user_id"] == student["device_user_id"]

    with pytest.raises(roster.RosterError, match="already belongs to ONE Person"):
        roster.edit_student(conn, student["id"], "ADEYEMI-COLE", "Direct Entry", "A/1", 1)
    with pytest.raises(roster.RosterError, match="surname"):
        roster.edit_student(conn, student["id"], "  ", "Direct Entry", None, 1)
    audit = conn.execute("SELECT detail FROM audit_log WHERE action = 'person.edit' ORDER BY id").fetchall()
    assert "matric number none to 25/PHARM01/120" in audit[0]["detail"]


# -- DE, the lecturer's mark for direct entry (2026-09-23) --------------------------------------------
#
# The lecturer's corrected list writes "DE" where a direct-entry student's matric number would go. It
# means "no number yet", so twenty-four of them are not twenty-four copies of one number: stored as
# nothing, recognised by name, shown as DE.

@pytest.mark.parametrize("mark", ["DE", "de", "D.E.", "D. E", "Direct entry", "direct-entry"])
def test_de_in_the_matric_column_means_no_number_yet(mark):
    from rollcall import names
    assert names.is_direct_entry(mark) and names.matric_or_none(mark) is None


@pytest.mark.parametrize("real", ["25/PHARM01/001", "DEX/1", "22/PHARM01/048"])
def test_a_real_matric_number_is_never_taken_for_de(real):
    from rollcall import names
    assert not names.is_direct_entry(real) and names.matric_or_none(real) == real


def test_several_students_marked_de_all_come_in(conn):
    lvl = level(conn, 200)["id"]
    plan = roster.build_plan(conn, lvl, 1, 1, roster.parse_csv(
        "matric_no,name\nA/1,ONE Person\nDE,FIRST Direct\nDE,SECOND Direct\nD.E.,THIRD Direct\n"))
    assert plan["new"] == 4, plan["issues"]
    assert not any("duplicate" in i for row in plan["issues"] for i in row["issues"])
    roster.commit(conn, pending(conn, lvl, "matric_no,name\nA/1,ONE Person\nDE,FIRST Direct\n"
                                           "DE,SECOND Direct\nD.E.,THIRD Direct\n"), 1)
    got = people(conn)
    assert [got[n]["matric_no"] for n in ("FIRST Direct", "SECOND Direct", "THIRD Direct")] == [None] * 3
    # Importing the same list again recognises them by name rather than making them twice.
    again = roster.build_plan(conn, lvl, 1, 1, roster.parse_csv(
        "matric_no,name\nA/1,ONE Person\nDE,FIRST Direct\nDE,SECOND Direct\nD.E.,THIRD Direct\n"))
    assert again["new"] == 0 and again["unchanged"] == 4


def test_typing_de_on_the_sheet_leaves_the_number_empty(conn):
    lvl = level(conn, 200)["id"]
    roster.commit(conn, pending(conn, lvl, "matric_no,name\nA/9,ADEYEMI Direct Entry\n"), 1)
    student = people(conn)["ADEYEMI Direct Entry"]
    roster.edit_student(conn, student["id"], "ADEYEMI", "Direct Entry", "DE", 1)
    assert people(conn)["ADEYEMI Direct Entry"]["matric_no"] is None
