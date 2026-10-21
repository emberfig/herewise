"""What a roster import does with the wrong file, a file in the wrong shape, and the wrong list.

Asked for on 2026-09-24, before the program went to a lecturer's laptop: "if a lecturer tries to
import an excel file or csv and its invalid what should we do about that". Each case here was fed
to the importer first and did something unhelpful or wrong; the probe is in
docs/DECISIONS-AND-LESSONS.md under that date. Every name here is made up.
"""
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from rollcall import db, roster
from rollcall.app import create_app

SIX = [("25/TEST01/%03d" % n, name) for n, name in enumerate(
    ["ADEYEMI Tolu Grace", "BELLO Samuel Ade", "CHUKWU Ada Nneka", "DANJUMA Musa Idris",
     "EZE Chioma Joy", "FASHOLA Kemi Ruth"], start=1)]


def csv_of(header, rows, sep=","):
    return (sep.join(header) + "\n" + "".join(sep.join(r) + "\n" for r in rows)).encode()


def workbook(sheets: dict) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def level(conn, number):
    return conn.execute("SELECT * FROM levels WHERE level_number = ?", (number,)).fetchone()["id"]


def plan_for(conn, raw, filename="roster.csv", number=200, **kwargs):
    rows = roster.parse_upload(raw, filename, number, **kwargs)
    return roster.build_plan(conn, level(conn, number), 1, 1, rows)


def commit(conn, raw, number=200, confirmed=(), filename="roster.csv"):
    rows = roster.parse_upload(raw, filename, number)
    import_id = conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, 1, 1, ?, 'pending', ?)", (level(conn, number), roster.as_csv(rows), db.now())).lastrowid
    return roster.commit(conn, import_id, None, confirmed)


def refusal(raw, filename="roster.csv"):
    with pytest.raises(roster.RosterError) as raised:
        roster.parse_upload(raw, filename, 200)
    assert not isinstance(raised.value, roster.NeedsColumns), "a refusal, not a question"
    return str(raised.value)


def zipped(**members) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, body in members.items():
            archive.writestr(name.replace("__", "/"), body)
    return out.getvalue()


# -- a file that is not a roster says what it is -------------------------------------------------

@pytest.mark.parametrize("raw, filename, said", [
    (b"%PDF-1.4\n%\xe2\xe3\n1 0 obj", "roster.xlsx", "roster.xlsx is a PDF, not a spreadsheet"),
    (b"\x89PNG\r\n\x1a\n" + bytes(64), "roster.csv", "is a picture (PNG)"),
    (b"\xff\xd8\xff\xe0" + bytes(64), "IMG_2044.jpg", "is a picture (JPEG)"),
    (zipped(word__document_xml="<w:document/>"), "roster.docx", "is a Word document"),
    (zipped(ppt__presentation_xml="<p/>"), "roster.pptx", "is a PowerPoint presentation"),
    (zipped(mimetype="application/vnd.oasis.opendocument.spreadsheet"), "roster.ods",
     "is a LibreOffice spreadsheet (.ods)"),
    (zipped(photos__a_jpg="x"), "rosters.zip", "is a zip folder"),
    (roster.XLS_MAGIC + bytes(512) + "WordDocument".encode("utf-16-le"), "old.doc", "is a Word document"),
    (b"MZ\x90\x00" + bytes(64), "setup.exe", "is a program"),
    (b"{\\rtf1\\ansi list}", "roster.rtf", "is a Rich Text document"),
    (b"\x00\x01\x02binary\x00stuff" * 20, "data.bin", "is not a spreadsheet or a CSV file"),
])
def test_a_file_that_is_not_a_roster_says_what_it_is(raw, filename, said):
    """A PDF used to be "not readable text, save it as CSV" and a Word document "could not be read
    as a spreadsheet": true, and no help to somebody who picked the wrong file."""
    assert said in refusal(raw, filename)


def test_a_workbook_locked_with_a_password_says_how_to_unlock_it():
    raw = roster.XLS_MAGIC + bytes(512) + "EncryptionInfo".encode("utf-16-le") + bytes(512)
    said = refusal(raw, "roster.xlsx")
    assert "locked with a password" in said and "Encrypt with Password" in said


def test_a_file_over_five_megabytes_is_turned_away_before_it_is_read():
    said = refusal(b"x" * (roster.MAX_UPLOAD_BYTES + 1), "video.csv")
    assert "video.csv is 5.0 MB" in said and "probably not the list" in said


def test_a_cell_longer_than_a_roster_is_refused_rather_than_crashing():
    """A 20 MB line of text stopped the reader with csv.Error, an error the screen could not word."""
    assert "longer than a whole roster" in refusal(b"x" * 200_000)


def test_headings_with_nobody_under_them_say_so():
    """They used to reach the preview, which said "Nothing in this file is new"."""
    said = refusal(b"matric_no,name\n")
    assert "no students under them" in said and "(matric_no, name) on row 1" in said


def test_an_unfilled_template_is_headings_with_nobody_under_them():
    assert "no students under them" in refusal(roster.template("200 Level"), "template.xlsx")


# -- files in a shape that is fine are read ------------------------------------------------------

TABBED = "matric_no\tname\n" + "".join(f"{m}\t{n}\n" for m, n in SIX)


@pytest.mark.parametrize("raw", [
    csv_of(["matric_no", "name"], SIX, "\t"),
    csv_of(["matric_no", "name"], SIX, ";"),
    TABBED.encode("utf-16"),
    TABBED.encode("utf-16-le"),
], ids=["tabs", "semicolons", "Excel Unicode Text", "UTF-16 with no mark"])
def test_a_list_separated_by_tabs_or_semicolons_or_saved_as_unicode_text_is_read(conn, raw):
    """Each was refused as "Missing column(s)", the Unicode one with headings of "m a t r i c"."""
    plan = plan_for(conn, raw, "roster.txt")
    assert (plan["new"], plan["skipped"]) == (6, 0)
    assert plan["students"][0]["surname"] == "ADEYEMI" and plan["students"][0]["other_names"] == "Tolu Grace"


def test_a_surname_in_a_column_of_its_own_is_kept_whole_through_the_commit(conn):
    """SURNAME and OTHER NAMES was refused as "Missing column(s): name". A two-word surname is the
    case that matters: joined and split again by guesswork, ADE BELLO would lose half of it."""
    raw = csv_of(["S/N", "MATRIC NO", "SURNAME", "OTHER NAMES"],
                 [("1", "25/TEST01/001", "ADE BELLO", "Samuel Tolu"), ("2", "25/TEST01/002", "Eze", "Chioma")])
    plan = plan_for(conn, raw)
    assert [(s["surname"], s["other_names"]) for s in plan["students"]] == [("ADE BELLO", "Samuel Tolu"), ("Eze", "Chioma")]
    commit(conn, raw)
    saved = conn.execute("SELECT surname, other_names FROM people ORDER BY device_user_id").fetchall()
    assert [tuple(r) for r in saved] == [("ADE BELLO", "Samuel Tolu"), ("Eze", "Chioma")]


@pytest.mark.parametrize("header, row", [
    (["Matric No", "Last Name", "First Name", "Middle Name"], ["25/TEST01/001", "OKAFOR", "Chinaza", "Ifeoma"]),
    (["MATRIC NO", "NAME", "OTHER NAMES"], ["25/TEST01/001", "OKAFOR", "Chinaza Ifeoma"]),
    (["Matriculation Number", "Family name", "Given names"], ["25/TEST01/001", "OKAFOR", "Chinaza Ifeoma"]),
])
def test_every_common_way_of_splitting_a_name_is_read(conn, header, row):
    plan = plan_for(conn, csv_of(header, [row]))
    assert (plan["students"][0]["surname"], plan["students"][0]["other_names"]) == ("OKAFOR", "Chinaza Ifeoma")


def test_the_one_sheet_holding_a_list_is_read_behind_a_cover_page(conn):
    raw = workbook({"Cover": [["COLLEGE OF PHARMACY"], ["2026/2027"]],
                    "Students": [["matric_no", "name"], *SIX]})
    assert plan_for(conn, raw, "groupings.xlsx")["new"] == 6


def test_a_sheet_with_headings_and_nobody_under_them_is_not_taken_for_the_list(conn):
    raw = workbook({"Sheet1": [["matric_no", "name"], *SIX], "Sheet2": [["matric_no", "name"]]})
    assert plan_for(conn, raw, "book.xlsx")["new"] == 6


def test_several_lists_and_none_named_for_the_level_asks_which_sheet(conn):
    raw = workbook({"Year Two": [["matric_no", "name"], *SIX[:2]], "Year Three": [["matric_no", "name"], *SIX[2:]]})
    with pytest.raises(roster.NeedsColumns) as asked:
        roster.parse_upload(raw, "years.xlsx", 200)
    shown = asked.value.payload()
    assert shown["sheets"] == ["Year Two", "Year Three"] and shown["sheet"] == "Year Two"
    assert "none of them is named for 200 Level" in shown["reason"]
    assert len(roster.parse_upload(raw, "years.xlsx", 200, sheet="Year Three")) == 4


# -- columns that hold the wrong thing are asked about, not imported -----------------------------

def test_columns_the_wrong_way_round_are_asked_about_with_the_swap_offered(conn):
    """Headings the right way round over values the wrong way round imported six students called
    "25/TEST01/001" without a word."""
    raw = csv_of(["matric_no", "name"], [(n, m) for m, n in SIX])
    with pytest.raises(roster.NeedsColumns) as asked:
        roster.parse_upload(raw, "roster.csv", 200)
    shown = asked.value.payload()
    assert "the wrong way round" in shown["reason"]
    assert shown["guess"]["matric_col"] == 1 and shown["guess"]["name_col"] == 0
    assert shown["guess"]["first_row"] == 2
    assert shown["rows"][0] == {"row": 1, "cells": ["matric_no", "name"]}
    choice = roster.Choice(matric=1, name=0, first_row=2)
    plan = plan_for(conn, raw, choice=choice)
    assert plan["new"] == 6 and plan["students"][0]["matric_no"] == "25/TEST01/001"


def test_headings_that_name_nothing_are_asked_about_with_a_guess_from_the_values(conn):
    raw = csv_of(["No.", "Reg", "Person", "Group"],
                 [(str(i), m, n, "A") for i, (m, n) in enumerate(SIX, 1)])
    with pytest.raises(roster.NeedsColumns) as asked:
        roster.parse_upload(raw, "roster.csv", 200)
    guess = asked.value.payload()["guess"]
    assert (guess["matric_col"], guess["name_col"], guess["first_row"]) == (1, 2, 2)


def test_a_title_block_longer_than_the_search_is_asked_about_from_the_right_row(conn):
    raw = b"COLLEGE OF PHARMACY\n" * 22 + csv_of(["matric_no", "name"], SIX)
    with pytest.raises(roster.NeedsColumns) as asked:
        roster.parse_upload(raw, "roster.csv", 200)
    shown = asked.value.payload()
    assert shown["guess"]["first_row"] == 24
    assert 24 in [r["row"] for r in shown["rows"]], "the first student is among the rows shown"


def test_chosen_columns_can_split_the_name_and_read_the_level(conn):
    raw = csv_of(["x", "y", "z", "w"], [(m, n.split()[0], " ".join(n.split()[1:]), "200") for m, n in SIX])
    plan = plan_for(conn, raw, choice=roster.Choice(matric=0, name=1, first_row=2, other=2, level=3))
    assert plan["new"] == 6
    assert (plan["students"][0]["surname"], plan["students"][0]["other_names"]) == ("ADEYEMI", "Tolu Grace")


def test_a_choice_that_finds_nothing_says_so():
    with pytest.raises(roster.RosterError, match="Nothing was found from row 9"):
        roster.parse_upload(csv_of(["a", "b"], SIX), "roster.csv", 200, choice=roster.Choice(0, 1, 9))


# -- rows that are not students ------------------------------------------------------------------

def test_a_total_under_the_list_is_refused_rather_than_imported(conn):
    """It has no matric number, so it used to go in as a student without one."""
    plan = plan_for(conn, csv_of(["matric_no", "name"], SIX + [("", "TOTAL 6")]))
    assert plan["new"] == 6 and plan["skipped"] == 1
    assert "the name has a number in it" in plan["issues"][0]["issues"][0]


def test_a_matric_number_with_no_number_in_it_is_refused_and_says_what_to_type(conn):
    plan = plan_for(conn, csv_of(["matric_no", "name"], SIX[:2] + [("N/A", "GANIYU Bola")]))
    assert plan["skipped"] == 1
    assert "type DE, or leave it empty" in plan["issues"][0]["issues"][0]


# -- the wrong list: importable, and probably not the one meant ----------------------------------

def test_another_levels_list_needs_a_tick_before_it_moves_them(conn):
    raw = csv_of(["matric_no", "name"], SIX)
    commit(conn, raw, number=300)
    plan = plan_for(conn, raw, number=200)
    assert [w["code"] for w in plan["warnings"]] == ["moves"]
    assert "6 of the 6 students in this file are on 300 Level already" in plan["warnings"][0]["text"]
    assert plan["warnings"][0]["confirm"] == "Yes, move them to 200 Level"
    with pytest.raises(roster.RosterError, match="Tick the box"):
        commit(conn, raw, number=200)
    assert commit(conn, raw, number=200, confirmed=["moves"])["updated"] == 6


def test_two_students_moving_level_is_not_warned_about(conn):
    """A corrected list with a couple of students on it who changed level is ordinary."""
    commit(conn, csv_of(["matric_no", "name"], SIX[:2]), number=300)
    assert plan_for(conn, csv_of(["matric_no", "name"], SIX))["warnings"] == []


def test_a_file_mostly_refused_needs_a_tick_before_the_rest_goes_in(conn):
    raw = csv_of(["matric_no", "name"], [(m, "") for m, _ in SIX[:4]] + SIX[4:])
    plan = plan_for(conn, raw)
    assert [w["code"] for w in plan["warnings"]] == ["mostly_refused"]
    assert plan["warnings"][0]["confirm"] == "Import the 2 that can be imported anyway"
    with pytest.raises(roster.RosterError, match="Tick the box"):
        commit(conn, raw)
    assert commit(conn, raw, confirmed=["mostly_refused"])["new"] == 2


def test_half_refused_is_not_mostly_refused(conn):
    raw = csv_of(["matric_no", "name"], [(m, "") for m, _ in SIX[:3]] + SIX[3:])
    assert plan_for(conn, raw)["warnings"] == []


# -- the template ----------------------------------------------------------------------------------

def test_the_template_is_named_for_the_level_and_reads_back_without_a_question(conn):
    book = load_workbook(io.BytesIO(roster.template("200 Level")))
    sheet = book.active
    assert sheet.title == "200 Level"
    assert [c.value for c in sheet[roster.TEMPLATE_HEADING_ROW]] == list(roster.TEMPLATE_HEADINGS)
    for matric, name in SIX[:2]:
        sheet.append([matric, name.split()[0], " ".join(name.split()[1:])])
    out = io.BytesIO()
    book.save(out)
    plan = plan_for(conn, out.getvalue(), "200 Level roster template.xlsx")
    assert plan["new"] == 2 and plan["students"][0]["surname"] == "ADEYEMI"


# -- through the API -------------------------------------------------------------------------------

@pytest.fixture
def client(tmp_path, k40_factory):
    with TestClient(create_app(tmp_path / "bad.db", k40_factory=k40_factory)) as c:
        c.post("/api/setup", json={"username": "lecturer", "display_name": "Dr", "password": "secret123",
                                   "academic_session": "2026/2027"})
        c.level = next(l["id"] for l in c.get("/api/levels").json() if l["level_number"] == 200)
        yield c


def test_the_preview_asks_which_column_is_which_then_imports_what_was_chosen(client):
    raw = csv_of(["matric_no", "name"], [(n, m) for m, n in SIX])
    asked = client.post(f"/api/levels/{client.level}/roster/preview", files={"file": ("r.csv", raw)})
    assert asked.status_code == 200 and asked.json()["needs_columns"] is True
    guess = asked.json()["guess"]
    chosen = client.post(f"/api/levels/{client.level}/roster/preview", files={"file": ("r.csv", raw)},
                         data={"matric_col": guess["matric_col"], "name_col": guess["name_col"],
                               "first_row": guess["first_row"]}).json()
    assert chosen["new"] == 6 and chosen["warnings"] == []
    done = client.post(f"/api/roster-imports/{chosen['import_id']}/commit", json={})
    assert done.status_code == 200, done.text
    students = client.get(f"/api/levels/{client.level}/students").json()["students"]
    assert sorted(s["matric_no"] for s in students) == [m for m, _ in SIX]


def test_the_commit_refuses_an_unticked_warning_and_takes_a_ticked_one(client):
    raw = csv_of(["matric_no", "name"], [(m, "") for m, _ in SIX[:4]] + SIX[4:])
    plan = client.post(f"/api/levels/{client.level}/roster/preview", files={"file": ("r.csv", raw)}).json()
    refused = client.post(f"/api/roster-imports/{plan['import_id']}/commit", json={})
    assert refused.status_code == 400 and "Tick the box" in refused.json()["error"]
    ticked = client.post(f"/api/roster-imports/{plan['import_id']}/commit", json={"confirm": ["mostly_refused"]})
    assert ticked.status_code == 200 and ticked.json()["new"] == 2


def test_a_commit_with_no_body_still_works_for_a_file_with_nothing_to_tick(client):
    plan = client.post(f"/api/levels/{client.level}/roster/preview",
                       files={"file": ("r.csv", csv_of(["matric_no", "name"], SIX))}).json()
    assert client.post(f"/api/roster-imports/{plan['import_id']}/commit").status_code == 200


def test_a_refusal_through_the_api_is_the_sentence(client):
    said = client.post(f"/api/levels/{client.level}/roster/preview",
                       files={"file": ("list.pdf", b"%PDF-1.7\n")})
    assert said.status_code == 400 and said.json()["error"].startswith("list.pdf is a PDF")


def test_a_huge_upload_is_refused_through_the_api(client):
    said = client.post(f"/api/levels/{client.level}/roster/preview",
                       files={"file": ("big.csv", b"x" * (roster.MAX_UPLOAD_BYTES + 10))})
    assert said.status_code == 400 and "probably not the list" in said.json()["error"]


def test_the_template_downloads_named_for_the_level(client):
    got = client.get(f"/api/levels/{client.level}/roster/template")
    assert got.status_code == 200
    assert got.headers["content-disposition"] == 'attachment; filename="200 Level roster template.xlsx"'
    assert load_workbook(io.BytesIO(got.content)).active.title == "200 Level"
