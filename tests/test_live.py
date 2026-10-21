"""A class taken live: Start in the room, the queue scans, End when it is done.

Asked for on 2026-09-22: "open an attendance session that I can start and end, it doesn't have to
be timed ... while students come to use their fingerprint it starts to pick them as present, and
it's only when I end that attendance session that it'll pick those students that have not signed as
absent ... (I can have more than one class for a particular level in a week)."
"""
from datetime import date, datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from conftest import ROSTER
from fake_zk import FakeRecord
from rollcall import db, live
from rollcall.app import create_app


@pytest.fixture(autouse=True)
def fresh_terminal_memory():
    """What the terminal said last is kept in memory per device; each test starts without it."""
    live._seen.clear()
    yield
    live._seen.clear()


@pytest.fixture
def client(tmp_path, k40_factory):
    app = create_app(tmp_path / "live.db", k40_factory=k40_factory)
    with TestClient(app) as c:
        c.db_path = tmp_path / "live.db"
        yield c


def a_course(client):
    client.post("/api/setup", json={"username": "lecturer", "display_name": "Dr Lecturer",
                                    "password": "secret123", "academic_session": "2026/2027"})
    level = next(l["id"] for l in client.get("/api/levels").json() if l["level_number"] == 200)
    preview = client.post(f"/api/levels/{level}/roster/preview",
                          files={"file": ("roster.csv", ROSTER.read_bytes(), "text/csv")})
    client.post(f"/api/roster-imports/{preview.json()['import_id']}/commit")
    client.post("/api/courses", json={"level_id": level, "code": "200 Level"})
    course = client.get("/api/courses").json()[0]
    students = client.get(f"/api/courses/{course['id']}/students").json()
    return course, students


def scan(fake, student, at: datetime):
    fake.records.append(FakeRecord(user_id=str(student["device_user_id"]), timestamp=at.replace(microsecond=0)))


def statuses(register):
    return {s["person_id"]: s["status"] for s in register["students"]}


def the_sheet(client, course):
    template_id = client.get(f"/api/courses/{course['id']}").json()["template_id"]
    return client.get(f"/api/templates/{template_id}/sheet").json()


def test_a_class_taken_live_marks_the_scanned_present_and_nobody_absent_until_it_ends(client, fake):
    course, students = a_course(client)
    started = client.post(f"/api/courses/{course['id']}/live")
    assert started.status_code == 200, started.text
    class_id = started.json()["register"]["session"]["id"]
    assert started.json()["register"]["open"] is True

    first, second = students[0], students[1]
    scan(fake, first, datetime.now())
    live._seen.clear()   # the next look is not throttled
    during = client.post(f"/api/live/{class_id}/refresh").json()
    seen = statuses(during["register"])
    assert seen[first["person_id"]] == "present"
    assert seen[second["person_id"]] == "waiting", "absent before the lecturer ended the class"
    assert during["register"]["totals"]["absent"] == 0
    assert during["terminal"]["state"] == "connected"

    ended = client.post(f"/api/live/{class_id}/end", json={})
    assert ended.status_code == 200, ended.text
    after = statuses(ended.json())
    assert after[first["person_id"]] == "present"
    assert after[second["person_id"]] == "absent", "nobody became absent when the class ended"
    assert ended.json()["open"] is False


def test_an_open_class_is_counted_nowhere(client, fake):
    """An open class with everyone not yet scanned would drag every percentage down for the length
    of a lecture, and put "Classes held 1" on a card before the class was done."""
    course, students = a_course(client)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]

    card = next(c for c in client.get("/api/courses").json() if c["id"] == course["id"])
    assert card["classes_held"] == 0
    sheet = the_sheet(client, course)
    assert sheet["classes_held"] == 0
    assert sheet["class_open"] is True
    open_slots = [s for w in sheet["weeks"] for s in w["slots"] if s["state"] == "open"]
    assert len(open_slots) == 1
    row = sheet["students"][0]
    assert row["held"] == 0 and row["percent"] == 0.0
    assert any(c.get("open") and c["state"] == "waiting" for c in row["cells"])

    client.post(f"/api/live/{class_id}/end", json={})
    sheet = the_sheet(client, course)
    assert sheet["classes_held"] == 1
    assert sheet["class_open"] is False


def test_only_scans_between_start_and_end_count(client, fake):
    """A class started and ended by hand needs no grace, and a grace would count a scan made for
    the class before it in the same room."""
    course, students = a_course(client)
    before_it, during_it = students[0], students[1]
    scan(fake, before_it, datetime.now() - timedelta(minutes=5))   # inside the usual 15 min grace
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    scan(fake, during_it, datetime.now())
    ended = client.post(f"/api/live/{class_id}/end", json={}).json()
    seen = statuses(ended)
    assert seen[during_it["person_id"]] == "present"
    assert seen[before_it["person_id"]] == "absent", "a scan from before Start was counted"


def test_one_class_takes_attendance_at_a_time(client, fake):
    course, _ = a_course(client)
    assert client.post(f"/api/courses/{course['id']}/live").status_code == 200
    again = client.post(f"/api/courses/{course['id']}/live")
    assert again.status_code == 400
    assert "already taking attendance" in again.json()["error"]


def test_starting_sets_the_terminal_clock_from_this_computer(client, fake):
    """Scans are stamped by the terminal and the window by this computer: a terminal five minutes
    slow would put the first five minutes of the queue outside the class."""
    course, _ = a_course(client)
    fake.clock_offset = timedelta(minutes=-5)
    client.post(f"/api/courses/{course['id']}/live")
    assert "set_time" in fake.calls
    assert abs(fake.clock_offset.total_seconds()) < 5


def test_a_class_is_not_started_on_a_terminal_that_does_not_answer(client, fake):
    course, _ = a_course(client)
    fake.reachable = False
    refused = client.post(f"/api/courses/{course['id']}/live")
    assert refused.status_code >= 400
    assert client.get("/api/live").json()["open"] is None


def test_a_loose_cable_during_the_class_leaves_it_open_and_says_so(client, fake):
    course, students = a_course(client)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    fake.reachable = False
    live._seen.clear()
    during = client.post(f"/api/live/{class_id}/refresh")
    assert during.status_code == 200
    assert during.json()["terminal"]["state"] == "unreachable"
    assert client.get("/api/live").json()["open"]["class_id"] == class_id

    # And the scans still count once the terminal answers again.
    fake.reachable = True
    scan(fake, students[0], datetime.now())
    ended = client.post(f"/api/live/{class_id}/end", json={}).json()
    assert statuses(ended)[students[0]["person_id"]] == "present"


def test_scans_read_during_the_class_are_one_line_in_the_activity_log(client, fake):
    """Every pull writes a line; a class read every few seconds would write one per student."""
    course, students = a_course(client)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    for student in students[:5]:
        scan(fake, student, datetime.now())
        live._seen.clear()
        client.post(f"/api/live/{class_id}/refresh")
    client.post(f"/api/live/{class_id}/end", json={})

    conn = db.connect(client.db_path)
    actions = [r["action"] for r in conn.execute("SELECT action FROM audit_log ORDER BY id")]
    batch = conn.execute("SELECT live_batch_id FROM class_sessions WHERE id = ?", (class_id,)).fetchone()[0]
    stored = conn.execute("SELECT COUNT(*) FROM punches WHERE import_batch_id = ?", (batch,)).fetchone()[0]
    conn.close()
    assert actions.count("punches.import.network") == 0
    assert actions.count("session.live_start") == 1
    assert actions.count("session.live_end") == 1
    assert stored == 5


def test_two_classes_in_one_week_both_go_on_the_sheet(client, fake):
    """A week set up with one box gains a second when a second class is taken in it, rather than
    leaving the class off the sheet with a warning."""
    course, students = a_course(client)
    for student in (students[0], students[1]):
        class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
        scan(fake, student, datetime.now())
        client.post(f"/api/live/{class_id}/end", json={})

    sheet = the_sheet(client, course)
    assert sheet["classes_held"] == 2
    this_week = next(w for w in sheet["weeks"] if w["starts_on"] <= date.today().isoformat() <= w["ends_on"])
    assert this_week["classes"] == 2
    assert all(s["state"] == "held" for s in this_week["slots"])
    assert not any("allows" in w for w in sheet["warnings"]), sheet["warnings"]


def test_the_first_class_sets_week_one_when_nobody_has(client, fake):
    """Set at Start, so the class is on the sheet, filling in, while the queue is still scanning."""
    course, _ = a_course(client)
    assert the_sheet(client, course)["weeks"][0]["starts_on"] is None
    client.post(f"/api/courses/{course['id']}/live")
    monday = date.today() - timedelta(days=date.today().weekday())
    assert the_sheet(client, course)["weeks"][0]["starts_on"] == monday.isoformat()


def test_a_class_ended_late_can_be_given_the_time_it_really_ended(client, fake):
    course, students = a_course(client)
    started_at = datetime.now() - timedelta(minutes=90)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    conn = db.connect(client.db_path)
    # As if Start was pressed an hour and a half ago.
    conn.execute("UPDATE class_sessions SET opened_at = ?, date = ?, start_time = ? WHERE id = ?",
                 (started_at.isoformat(timespec="seconds"), started_at.date().isoformat(),
                  started_at.strftime("%H:%M"), class_id))
    conn.close()
    in_class, after_class = students[0], students[1]
    scan(fake, in_class, started_at + timedelta(minutes=10))
    scan(fake, after_class, datetime.now() - timedelta(minutes=5))
    lecture_ended = (started_at + timedelta(minutes=60)).strftime("%H:%M")

    too_early = client.post(f"/api/live/{class_id}/end",
                            json={"ended_at": (started_at - timedelta(minutes=5)).strftime("%H:%M")})
    assert too_early.status_code == 400
    ended = client.post(f"/api/live/{class_id}/end", json={"ended_at": lecture_ended}).json()
    seen = statuses(ended)
    assert seen[in_class["person_id"]] == "present"
    assert seen[after_class["person_id"]] == "absent", "a scan after the given end was counted"


def test_enrolment_waits_while_a_class_is_taking_attendance(client, fake):
    course, students = a_course(client)
    client.post(f"/api/courses/{course['id']}/live")
    refused = client.post("/api/enrollment/start", json={"person_id": students[0]["person_id"],
                                                         "finger_index": 6, "consent_confirmed": True})
    assert refused.status_code == 409
    assert "taking attendance" in refused.json()["error"]


def test_a_pull_during_a_long_class_does_not_make_a_second_class_of_it(client, fake):
    """Detection rebuilt a class's window from its times plus the grace, which for a class still
    open ends half an hour after it began: a pull an hour into a lecture would have made a second
    class out of everyone who scanned late."""
    course, students = a_course(client)
    started_at = datetime.now() - timedelta(minutes=70)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    conn = db.connect(client.db_path)
    conn.execute("UPDATE class_sessions SET opened_at = ?, date = ?, start_time = ?, end_time = ?"
                 " WHERE id = ?",
                 (started_at.isoformat(timespec="seconds"), started_at.date().isoformat(),
                  started_at.strftime("%H:%M"), started_at.strftime("%H:%M"), class_id))
    conn.close()
    for student in students[:8]:
        scan(fake, student, datetime.now() - timedelta(minutes=2))
    client.post("/api/devices/1/pull", json={})

    conn = db.connect(client.db_path)
    classes = conn.execute("SELECT COUNT(*) FROM class_sessions").fetchone()[0]
    conn.close()
    assert classes == 1, "a pull made a second class out of scans inside the open one"


def test_the_sheet_and_the_course_header_quote_the_same_planned_classes(client, fake):
    """With a saved plan of 12, a week that gained a box for a second class made the sheet say
    "7 of 13 classes held" under a header saying "7 of 12" (seen on the screen, 2026-09-23)."""
    course, students = a_course(client)
    saved = client.put(f"/api/courses/{course['id']}/attendance-plan", json={"total_classes": 12})
    assert saved.status_code == 200, saved.text
    for student in (students[0], students[1]):
        class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
        scan(fake, student, datetime.now())
        client.post(f"/api/live/{class_id}/end", json={})

    sheet = the_sheet(client, course)
    header = client.get(f"/api/courses/{course['id']}?with_figures=true").json()["figures"]
    assert sum(w["classes"] for w in sheet["weeks"]) == 13, "the week did not gain its box"
    assert sheet["planned_classes"] == header["planned_classes"] == 12


def test_an_export_made_during_the_class_marks_nobody_absent(client, fake):
    """A student who has not scanned yet is not absent. Exported mid-class, the waiting boxes were
    crosses: an absence on paper that the register does not hold."""
    import io

    from openpyxl import load_workbook

    from rollcall import sheet_export, templates

    course, students = a_course(client)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    scan(fake, students[0], datetime.now())
    live._seen.clear()
    client.post(f"/api/live/{class_id}/refresh")

    conn = db.connect(client.db_path)
    template_id = templates.ensure_template(conn, course["id"])
    sheet = templates.grid(conn, template_id)
    conn.close()
    ws = load_workbook(io.BytesIO(sheet_export.to_xlsx(sheet))).active
    marks = [ws.cell(row, 4).value for row in range(5, 5 + len(students))]
    assert marks.count(sheet_export.TICK) == 1
    assert sheet_export.CROSS not in marks, "a student who had not scanned yet was exported as absent"
    assert sheet_export.to_pdf(sheet)[:5] == b"%PDF-"


# -- which week a class goes in ------------------------------------------------------------------
#
# Asked for on 2026-09-23: "sometimes we can have more than one class in a week so when adding the
# collected attendance the terminal should ask if the attendance is in the first week or the second
# week, and continually do that with the other weeks."


def take(client, fake, course, student, week=None):
    body = {"week_number": week} if week is not None else {}
    started = client.post(f"/api/courses/{course['id']}/live", json=body)
    assert started.status_code == 200, started.text
    class_id = started.json()["register"]["session"]["id"]
    if student is not None:
        scan(fake, student, datetime.now())
    return client.post(f"/api/live/{class_id}/end", json={}).json()


def held_in(sheet, label):
    week = next(w for w in sheet["weeks"] if w["label"] == label)
    return [s["state"] for s in week["slots"]]


def test_the_week_chosen_is_where_the_class_goes_whatever_the_date(client, fake):
    course, students = a_course(client)
    first = take(client, fake, course, students[0], week=1)
    assert first["sheet"]["week"] == "WK 1" and first["sheet"]["position"] == 1
    # Taken today, but the lecturer says it is week 3: a make-up, or the weeks are not the calendar's.
    third = take(client, fake, course, students[1], week=3)
    assert third["sheet"]["week"] == "WK 3"
    sheet = the_sheet(client, course)
    assert held_in(sheet, "WK 1") == ["held"]
    assert held_in(sheet, "WK 3") == ["held"]
    assert held_in(sheet, "WK 2")[0] != "held"


def test_a_second_class_in_the_same_week_is_its_second_box(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=2)
    second = take(client, fake, course, students[1], week=2)
    assert second["sheet"]["week"] == "WK 2" and second["sheet"]["position"] == 2
    sheet = the_sheet(client, course)
    assert held_in(sheet, "WK 2") == ["held", "held"]
    assert not sheet["warnings"], sheet["warnings"]


def test_the_first_class_in_week_four_sets_week_one_three_weeks_earlier(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=4)
    monday = date.today() - timedelta(days=date.today().weekday())
    sheet = the_sheet(client, course)
    assert sheet["weeks"][0]["starts_on"] == (monday - timedelta(weeks=3)).isoformat()
    assert held_in(sheet, "WK 4") == ["held"]


def test_the_weeks_offered_say_what_each_holds_and_suggest_one(client, fake):
    course, students = a_course(client)
    fresh = client.get(f"/api/courses/{course['id']}/live/weeks").json()
    assert fresh["suggested"] == 1, "a course with nothing taken starts at week 1"
    assert len(fresh["weeks"]) == 12 and all(w["held"] == 0 for w in fresh["weeks"])

    take(client, fake, course, students[0], week=1)
    after = client.get(f"/api/courses/{course['id']}/live/weeks").json()
    week1 = after["weeks"][0]
    assert week1["held"] == 1 and week1["current"] is True
    assert after["suggested"] == 1, "the week today falls in, once the weeks have dates"


def test_a_week_the_sheet_does_not_have_is_refused(client, fake):
    course, _ = a_course(client)
    refused = client.post(f"/api/courses/{course['id']}/live", json={"week_number": 99})
    assert refused.status_code == 400
    assert "Sheet setup" in refused.json()["error"]
    assert client.get("/api/live").json()["open"] is None


def test_the_class_taking_attendance_says_its_week_and_place(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=5)
    client.post(f"/api/courses/{course['id']}/live", json={"week_number": 5})
    open_now = client.get("/api/live").json()["open"]
    assert (open_now["week"], open_now["position"], open_now["of"]) == ("WK 5", 2, 2)


def test_out_of_term_the_suggestion_is_the_week_after_the_last_one_used(client, fake):
    """No week holds today (a make-up in the holiday, or dates never set): the next class usually
    goes in the week after the last one that holds a class."""
    course, students = a_course(client)
    take(client, fake, course, students[0], week=1)
    take(client, fake, course, students[1], week=4)
    conn = db.connect(client.db_path)
    far = live.week_choices(conn, course["id"], today=date.today() + timedelta(days=400))
    conn.close()
    assert not any(w["current"] for w in far["weeks"])
    assert far["suggested"] == 5


def test_a_class_put_in_the_wrong_week_can_be_moved(client, fake):
    """Without a way back, a class started in the wrong week stayed there."""
    course, students = a_course(client)
    take(client, fake, course, students[0], week=2)
    class_id = client.get(f"/api/sessions?course_id={course['id']}").json()[0]["id"]
    moved = client.put(f"/api/sessions/{class_id}/week", json={"week_number": 3})
    assert moved.status_code == 200, moved.text
    assert moved.json()["week"] == "WK 3"
    sheet = the_sheet(client, course)
    assert held_in(sheet, "WK 3") == ["held"]
    assert held_in(sheet, "WK 2")[0] != "held"
    listed = client.get(f"/api/sessions?course_id={course['id']}").json()[0]
    assert listed["sheet_week"]["label"] == "WK 3", "the class list and the sheet disagree"
    assert client.put(f"/api/sessions/{class_id}/week", json={"week_number": 99}).status_code == 400


# -- weeks a lecturer can remove and add ---------------------------------------------------------
#
# Asked for on 2026-09-23: "lecturers don't know the amount of weeks they have so make weeks
# deletable".


def template_of(client, course):
    return client.get(f"/api/courses/{course['id']}/template").json()


def save_weeks(client, template, weeks):
    return client.put(f"/api/templates/{template['id']}/weeks", json={"weeks": weeks})


def as_sent(week):
    """A week as Sheet setup sends it back: its shown label included."""
    return {"classes": week["classes"], "starts_on": week["starts_on"], "label": week["label"],
            "from_week": week["week_number"]}


def test_removing_an_empty_week_moves_the_later_weeks_classes_with_them(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=7)
    t = template_of(client, course)
    assert t["weeks"][6]["held"] == 1 and t["weeks"][4]["held"] == 0
    kept = [as_sent(w) for w in t["weeks"] if w["week_number"] != 5]
    saved = save_weeks(client, t, kept)
    assert saved.status_code == 200, saved.text
    sheet = the_sheet(client, course)
    assert len(sheet["weeks"]) == 11
    assert held_in(sheet, "WK 6") == ["held"], "the class did not follow its week"
    assert all(w["label"] == f"WK {w['week_number']}" for w in sheet["weeks"]), "labels kept old numbers"


def test_a_week_holding_a_class_is_not_removed(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=3)
    t = template_of(client, course)
    refused = save_weeks(client, t, [as_sent(w) for w in t["weeks"] if w["week_number"] != 3])
    assert refused.status_code == 400
    assert "WK 3 holds 1 class" in refused.json()["error"]
    assert len(template_of(client, course)["weeks"]) == 12, "something was removed anyway"


def test_generating_fewer_weeks_cannot_drop_a_class(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=10)
    t = template_of(client, course)
    refused = client.post(f"/api/templates/{t['id']}/weeks/generate", json={"weeks": 8, "classes_per_week": 1})
    assert refused.status_code == 400
    assert "WK 10" in refused.json()["error"]


def test_trailing_empty_weeks_can_go_and_the_plan_follows(client, fake):
    course, students = a_course(client)
    take(client, fake, course, students[0], week=1)
    t = template_of(client, course)
    saved = save_weeks(client, t, [as_sent(w) for w in t["weeks"][:6]])
    assert saved.status_code == 200, saved.text
    assert the_sheet(client, course)["planned_classes"] == 6


def test_a_new_week_can_be_chosen_when_the_course_runs_longer(client, fake):
    course, students = a_course(client)
    choices = client.get(f"/api/courses/{course['id']}/live/weeks").json()
    assert choices["new_week"] == 13
    ended = take(client, fake, course, students[0], week=13)
    assert ended["sheet"]["week"] == "WK 13"
    sheet = the_sheet(client, course)
    assert len(sheet["weeks"]) == 13 and held_in(sheet, "WK 13") == ["held"]


def test_a_renamed_student_on_the_terminal_is_sent_the_new_name(client, fake):
    """The terminal shows the name at every scan, so an edit reaches it straight away."""
    course, students = a_course(client)
    student = students[0]
    pushed = client.post("/api/devices/1/push", json={"person_ids": [student["person_id"]]})
    assert pushed.status_code == 200, pushed.text
    edited = client.put(f"/api/people/{student['person_id']}",
                        json={"surname": "ADEBAYO-COLE", "other_names": "Takme", "matric_no": "25/X/1"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["terminal"] == "sent"
    assert fake.users[student["device_user_id"]].name.startswith("ADEBAYO-COLE")
    sheet = the_sheet(client, course)
    row = next(r for r in sheet["students"] if r["person_id"] == student["person_id"])
    assert (row["full_name"], row["matric_no"]) == ("ADEBAYO-COLE Takme", "25/X/1")


def test_editing_a_student_needs_the_roster_permission(client, fake):
    course, students = a_course(client)
    client.post("/api/accounts", json={"username": "viewer", "display_name": "Viewer", "password": "secret123",
                                       "role_id": next(r["id"] for r in client.get("/api/roles").json()
                                                       if r["name"] == "Viewer")})
    client.post("/api/logout")
    client.post("/api/login", json={"username": "viewer", "password": "secret123"})
    refused = client.put(f"/api/people/{students[0]['person_id']}",
                         json={"surname": "X", "other_names": "", "matric_no": None})
    assert refused.status_code == 403


# -- the terminal's clock (found on the owner's K40, 2026-09-23) ---------------------------------------
#
# Three starts running set the clock 9 and 10 seconds slow, because "now" was worked out before the
# wait for the terminal and written after it. Every scan was stamped that much early, and the first
# scan after Start fell before the class opened: 1 of 120 where the lecturer had scanned twice.

def test_a_terminal_slow_to_answer_is_still_set_to_the_right_time(client, fake):
    course, _ = a_course(client)
    fake.connect_delay = 1.5
    started = client.post(f"/api/courses/{course['id']}/live")
    assert started.status_code == 200, started.text
    assert abs(fake.clock_offset.total_seconds()) < 1, fake.clock_offset


def test_the_first_scan_after_start_counts_even_on_a_slow_clock(client, fake):
    """A terminal whose clock will not be set, ten seconds slow: the class opens on its clock."""
    course, students = a_course(client)
    fake.clock_offset = timedelta(seconds=-10)
    fake.clock_stuck = True
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    scan(fake, students[0], datetime.now() + fake.clock_offset)      # straight after Start
    live._seen.clear()   # the next look is not throttled
    during = client.post(f"/api/live/{class_id}/refresh").json()
    assert statuses(during["register"])[students[0]["person_id"]] == "present"


def test_cancelling_a_class_taking_attendance_stops_it(client, fake):
    """Reported 2026-09-23: cancelled, but the banner went on saying it was taking attendance."""
    course, students = a_course(client)
    class_id = client.post(f"/api/courses/{course['id']}/live").json()["register"]["session"]["id"]
    assert client.get("/api/live").json()["open"]["class_id"] == class_id
    cancelled = client.post(f"/api/sessions/{class_id}/cancel", json={"note": "Started by mistake."})
    assert cancelled.status_code == 200, cancelled.text
    assert client.get("/api/live").json()["open"] is None, "nothing is taking attendance any more"
    register = client.get(f"/api/sessions/{class_id}/register").json()
    assert register["open"] is False and register["session"]["cancelled_at"]
    # And the terminal is free for the next class, without anybody pressing End.
    assert client.post(f"/api/courses/{course['id']}/live").status_code == 200
