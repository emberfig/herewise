import io
import os
import re
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from conftest import ROSTER
from fake_zk import FakeRecord
from rollcall import courses, db, register, roster, sheet_export, templates
from rollcall.app import create_app
from rollcall.device import RawPunch
from rollcall.punches import store_punches

TICK, CROSS = sheet_export.TICK, sheet_export.CROSS


def import_level(conn, text=None):
    text = text or ROSTER.read_text(encoding="utf-8-sig")
    lvl = conn.execute("SELECT id FROM levels WHERE level_number = 200").fetchone()["id"]
    import_id = conn.execute(
        "INSERT INTO roster_imports(level_id, academic_session_id, device_id, csv_text, status, created_at)"
        " VALUES (?, 1, 1, ?, 'pending', ?)", (lvl, text, db.now())).lastrowid
    roster.commit(conn, import_id, 1)
    # An import makes no course (2026-09-24), so the first call makes the one these tests use.
    made = conn.execute("SELECT id FROM courses WHERE level_id = ?", (lvl,)).fetchone()
    course = made["id"] if made else courses.create_course(conn, 1, lvl, "200 Level", None, 1)
    return course, templates.ensure_template(conn, course)


def scan(conn, ids, when):
    store_punches(conn, 1, [RawPunch(str(i), datetime.fromisoformat(when), 1, 255) for i in ids], "network", 1)


def cell(sheet, device_user_id, week, slot):
    row = next(s for s in sheet["students"] if s["device_user_id"] == device_user_id)
    return next(c for c in row["cells"] if c["week"] == week and c["slot"] == slot)


def test_a_new_course_gets_the_default_template(conn):
    _, template_id = import_level(conn)
    t = templates.get_template(conn, template_id)
    assert t["title"] == "200 LEVEL ATTENDANCE"
    assert len(t["weeks"]) == 12 and all(w["classes"] == 1 for w in t["weeks"])
    assert [w["label"] for w in t["weeks"]][:2] == ["WK 1", "WK 2"]
    # Importing more names into the level keeps the same template.
    _, again = import_level(conn, "matric_no,name,level\nNEW/1,NEWSTUDENT Person,200\n")
    assert again == template_id


def test_weeks_and_classes_are_editable(conn):
    _, tid = import_level(conn)
    t = templates.generate_weeks(conn, tid, 1, weeks=12, classes_per_week=2, week1_start="2026-09-14")
    assert t["total_classes"] == 24 and t["weeks"][1]["starts_on"] == "2026-09-21"

    weeks = [{"classes": 2} for _ in range(13)]
    weeks[2]["classes"] = 0                            # no class in week 3
    weeks[6]["starts_on"] = "2026-11-09"               # two week break before week 7
    weeks[12]["label"] = "REVISION"
    t = templates.set_weeks(conn, tid, 1, weeks)
    assert len(t["weeks"]) == 13 and t["total_classes"] == 24
    assert t["weeks"][5]["starts_on"] == "2026-10-19"
    assert t["weeks"][6]["starts_on"] == "2026-11-09" and t["weeks"][7]["starts_on"] == "2026-11-16"
    assert t["weeks"][12]["label"] == "REVISION"

    # Moving week 1 moves everything and keeps the break.
    t = templates.update_template(conn, tid, 1, title="200 LEVEL PHARMACY ATTENDANCE", course_code="PCT 201",
                                  week1_start="2026-09-21")
    assert t["weeks"][6]["starts_on"] == "2026-11-16" and t["title"] == "200 LEVEL PHARMACY ATTENDANCE"

    with pytest.raises(templates.TemplateError, match="before week 1 has finished"):
        templates.set_weeks(conn, tid, 1, [{"classes": 1, "starts_on": "2026-09-14"},
                                           {"classes": 1, "starts_on": "2026-09-17"}])
    with pytest.raises(templates.TemplateError):
        templates.set_weeks(conn, tid, 1, [{"classes": 8}])


def test_scans_tick_the_right_box(conn):
    lvl, tid = import_level(conn)
    templates.generate_weeks(conn, tid, 1, weeks=12, classes_per_week=2, week1_start="2026-09-14")
    s1 = register.create_session(conn, lvl, "2026-09-15", "09:00", "11:00", "device", 1)
    s2 = register.create_session(conn, lvl, "2026-09-17", "14:00", "16:00", "device", 1)
    s3 = register.create_session(conn, lvl, "2026-09-18", "09:00", "10:00", "device", 1)   # one too many
    scan(conn, [2001, 2002], "2026-09-15T10:55:00")
    scan(conn, [2002], "2026-09-17T15:58:00")
    register.add_mark(conn, s2, conn.execute("SELECT id FROM people WHERE device_user_id = 2003").fetchone()[0],
                      "present", "Scanner queue", 1)

    sheet = templates.grid(conn, tid, today=date(2026, 9, 30))
    week1 = sheet["weeks"][0]
    assert [s["session_id"] for s in week1["slots"]] == [s1, s2]
    assert any("WK 1 has 3 classes" in w for w in sheet["warnings"]) and s3
    assert cell(sheet, 2001, 1, 1)["ticked"] and cell(sheet, 2001, 1, 1)["state"] == "present"
    assert cell(sheet, 2001, 1, 2) == {"week": 1, "slot": 2, "state": "absent", "ticked": False, "first_scan": None,
                                       "excusal_category": None}
    assert cell(sheet, 2003, 1, 2)["state"] == "marked_present" and cell(sheet, 2003, 1, 2)["ticked"]
    assert cell(sheet, 2001, 2, 1)["state"] == "not_held"          # week 2 is over, no class recorded
    assert cell(sheet, 2001, 5, 1)["state"] == "upcoming"
    assert week1["slots"][0]["present"] == 2 and week1["slots"][0]["absent"] == 256
    row = next(s for s in sheet["students"] if s["device_user_id"] == 2002)
    assert (row["sn"], row["attended"], row["held"]) == (2, 2, 2)
    assert sheet["total_registered"] == 258


def test_the_sheet_states_each_box_and_the_current_week_itself(conn):
    lvl, tid = import_level(conn)
    templates.generate_weeks(conn, tid, 1, weeks=12, classes_per_week=2, week1_start="2026-09-14")
    s1 = register.create_session(conn, lvl, "2026-09-15", "09:00", "11:00", "device", 1)
    bola = conn.execute("SELECT id FROM people WHERE device_user_id = 2004").fetchone()[0]
    register.add_excusal(conn, bola, "2026-09-15", "2026-09-15", "medical", "Clinic note 114", 1)

    sheet = templates.grid(conn, tid, today=date(2026, 9, 30))
    states = [[s["state"] for s in w["slots"]] for w in sheet["weeks"][:4]]
    assert states == [["held", "not_held"], ["not_held", "not_held"], ["upcoming", "upcoming"],
                      ["upcoming", "upcoming"]]
    assert sheet["weeks"][0]["slots"][0]["session_id"] == s1
    # 30 September falls in week 3 (28 September to 4 October), and in no other.
    assert [w["week_number"] for w in sheet["weeks"] if w["current"]] == [3]
    # Every student's box agrees with the box's own state.
    for row in sheet["students"]:
        assert [c["state"] for c in row["cells"][1:4]] == ["not_held", "not_held", "not_held"]
    assert cell(sheet, 2004, 1, 1)["state"] == "excused"
    assert cell(sheet, 2004, 1, 1)["excusal_category"] == "medical"


def test_no_start_date_warns_and_ticks_nothing(conn):
    lvl, tid = import_level(conn)
    register.create_session(conn, lvl, "2026-09-15", "09:00", "11:00", "device", 1)
    sheet = templates.grid(conn, tid)
    assert any("week 1 start date" in w for w in sheet["warnings"])
    assert sheet["classes_held"] == 0


def test_detects_classes_from_scans(conn):
    lvl, _ = import_level(conn)
    scan(conn, [2001, 2002, 2003], "2026-09-15T10:50:00")
    scan(conn, [2004, 2005, 2006], "2026-09-15T10:58:00")
    scan(conn, [2007, 2008], "2026-09-15T15:00:00")                 # too few for a class
    created = register.detect_sessions(conn, 1)
    assert [(c["date"], c["start_time"], c["end_time"], c["students"]) for c in created] == \
        [("2026-09-15", "10:50", "10:59", 6)]
    assert register.detect_sessions(conn, 1) == []                    # no duplicate on the next pull

    register.cancel_session(conn, created[0]["id"], "Not a lecture", 1)
    assert register.detect_sessions(conn, 1) == []                    # a cancel is respected

    scan(conn, [2010, 2011, 2012, 2013, 2014], "2026-09-16T09:00:00")
    scan(conn, [2015, 2016, 2017, 2018, 2019], "2026-09-16T13:00:00")  # separate class later that day
    assert len(register.detect_sessions(conn, 1)) == 2


def test_summary_uses_template_class_count(conn):
    lvl, tid = import_level(conn)
    templates.generate_weeks(conn, tid, 1, weeks=13, classes_per_week=2, week1_start="2026-09-14")
    summary = register.course_summary(conn, lvl)
    assert summary["planned_classes"] == 26 and summary["plan"]["planned_source"] == "template"


def sample_sheet(conn):
    lvl, tid = import_level(conn)
    templates.generate_weeks(conn, tid, 1, weeks=12, classes_per_week=1, week1_start="2026-09-14")
    templates.update_template(conn, tid, 1, title="200 LEVEL PHARMACY ATTENDANCE", course_code="PCT 201")
    register.create_session(conn, lvl, "2026-09-15", "09:00", "11:00", "device", 1)
    scan(conn, [2001, 2003], "2026-09-15T10:55:00")
    return templates.grid(conn, tid, today=date(2026, 9, 16))


def test_excel_export_matches_the_sheet(conn):
    from openpyxl import load_workbook

    sheet = sample_sheet(conn)
    ws = load_workbook(io.BytesIO(sheet_export.to_xlsx(sheet))).active
    assert ws["A1"].value == "200 LEVEL PHARMACY ATTENDANCE"
    assert ws["A2"].value == "COURSE: PCT 201"
    assert ws["A3"].value == "LECTURER'S SIGNATURE"
    assert [ws.cell(4, c).value for c in (1, 2, 3, 4, 15)] == ["S/N", "MATRIC NO.", "FULL NAME", "WK 1", "WK 12"]
    assert [ws.cell(5, c).value for c in (1, 2, 3)] == [1, "23/MHX01/090", "EFIKOSLI Opazsitofo Erimalsilke"]
    assert ws.cell(5, 4).value == TICK and ws.cell(6, 4).value == CROSS and ws.cell(5, 5).value is None
    assert ws.cell(263, 1).value == "NO. OF STUDENTS PRESENT"
    assert ws.cell(263, 4).value == f'=COUNTIF(D5:D262,"{TICK}")'
    assert ws.cell(265, 15).value == "TOTAL REGISTERED: 258"
    assert ws.page_setup.orientation == "landscape"


def test_pdf_export_has_seven_pages_for_258_students(conn):
    content = sheet_export.to_pdf(sample_sheet(conn))
    assert content[:5] == b"%PDF-"
    assert len(re.findall(rb"/Type /Page[^s]", content)) == 7


def page_boxes(content: bytes) -> set[tuple[int, int]]:
    return {(round(float(w)), round(float(h))) for w, h in re.findall(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", content)}


def test_pdf_paper_follows_the_setting_and_keeps_seven_pages(conn):
    sheet = sample_sheet(conn)
    a4, letter = sheet_export.to_pdf(sheet, "a4"), sheet_export.to_pdf(sheet, "letter")
    assert page_boxes(a4) == {(842, 595)} and page_boxes(letter) == {(792, 612)}
    assert len(re.findall(rb"/Type /Page[^s]", a4)) == len(re.findall(rb"/Type /Page[^s]", letter)) == 7


@pytest.mark.skipif(not (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "segoeui.ttf").is_file()
                    and not (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf").is_file(),
                    reason="needs Segoe UI or Arial from Windows")
def test_yoruba_and_igbo_names_reach_the_pdf_as_letters(conn):
    # Helvetica drew these as ZapfDingbats black squares: "Ọlábísí Ṣẹ́gun" came out "■lábísí ■■■gun".
    import_level(conn, "matric_no,name,level\nY/1,OLÁBÍSÍ Ọlábísí Ṣẹ́gun,200\nI/1,OKORO Chụkwụemeka Ọ̀kọ̀rọ̀,200\n")
    tid = conn.execute("SELECT id FROM register_templates").fetchone()["id"]
    content = sheet_export.to_pdf(templates.grid(conn, tid), "a4")
    assert b"ZapfDingbats" not in content
    assert b"/FontFile2" in content                    # an embedded TrueType font, not Helvetica


def test_pdf_text_keeps_dotted_letters_and_drops_only_accents_that_cannot_join():
    from rollcall.pdf_fonts import text

    assert text("Ọla") == "Ọla"                        # base plus dot below composes
    assert text("Ṣẹ́gun") == "Ṣẹgun"                          # no composed ẹ́ exists: the accent goes
    assert text("José Müller") == "José Müller"


def test_template_api_flow(tmp_path, k40_factory, fake):
    with TestClient(create_app(tmp_path / "t.db", k40_factory=k40_factory)) as c:
        c.post("/api/setup", json={"username": "lecturer", "display_name": "Dr", "password": "secret123",
                                   "academic_session": "2026/2027"})
        lvl = next(l["id"] for l in c.get("/api/levels").json() if l["level_number"] == 200)
        preview = c.post(f"/api/levels/{lvl}/roster/preview", files={"file": ("r.csv", ROSTER.read_bytes())}).json()
        c.post(f"/api/roster-imports/{preview['import_id']}/commit")
        course = c.post("/api/courses", json={"level_id": lvl, "code": "200 Level"}).json()["id"]
        tid = c.get(f"/api/courses/{course}/template").json()["id"]

        assert c.patch(f"/api/templates/{tid}", json={"title": "200 LEVEL PHARMACY ATTENDANCE",
                                                      "week1_start": "2026-09-14"}).status_code == 200
        weeks = c.put(f"/api/templates/{tid}/weeks",
                      json={"weeks": [{"classes": 2}] + [{"classes": 1}] * 11}).json()
        assert weeks["total_classes"] == 13

        fake.records = [FakeRecord(str(2001 + i), datetime(2026, 9, 15, 10, 50 + i)) for i in range(6)]
        pulled = c.post("/api/devices/1/pull").json()
        assert len(pulled["detected_sessions"]) == 1

        sheet = c.get(f"/api/templates/{tid}/sheet").json()
        assert sheet["weeks"][0]["slots"][0]["present"] == 6
        assert cell(sheet, 2006, 1, 1)["ticked"] and not cell(sheet, 2007, 1, 1)["ticked"]

        xlsx = c.get(f"/api/templates/{tid}/export", params={"format": "xlsx"})
        assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
        assert "200 LEVEL PHARMACY ATTENDANCE 2026-2027.xlsx" in xlsx.headers["content-disposition"]
        pdf = c.get(f"/api/templates/{tid}/export", params={"format": "pdf"})
        assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
        assert c.put(f"/api/templates/{tid}/weeks", json={"weeks": [{"classes": 9}]}).status_code == 422


def test_pages_leave_room_for_the_totals():
    assert sheet_export.paginate(258, False) == [(i, i + 38) for i in range(0, 228, 38)] + [(228, 258)]
    multi = sheet_export.paginate(258, True)
    assert all(b - a <= 37 for a, b in multi) and multi[-1][1] == 258
    assert multi[-1][1] - multi[-1][0] <= 34          # three rows free for present, absent, total
    assert sheet_export.paginate(37, False) == [(0, 35), (35, 37)]
    assert sheet_export.paginate(0, False) == [(0, 0)]


def test_a_direct_entry_student_reads_de_on_the_exported_sheet(tmp_path, k40_factory):
    """The lecturer's mark for a student with no matric number yet, on paper as on screen (2026-09-23)."""
    import io

    import openpyxl
    from fastapi.testclient import TestClient

    from rollcall.app import create_app

    app = create_app(tmp_path / "de.db", k40_factory=k40_factory)
    with TestClient(app) as c:
        c.post("/api/setup", json={"username": "lecturer", "display_name": "Dr Lecturer",
                                   "password": "secret123", "academic_session": "2026/2027"})
        lvl = next(l["id"] for l in c.get("/api/levels").json() if l["level_number"] == 200)
        roster_csv = "matric_no,name\n25/PHARM01/001,ONE Person\nDE,OKEBEGASO Tizaluziz Vovnuzf\n"
        preview = c.post(f"/api/levels/{lvl}/roster/preview",
                         files={"file": ("r.csv", roster_csv.encode())}).json()
        c.post(f"/api/roster-imports/{preview['import_id']}/commit")
        course = c.post("/api/courses", json={"level_id": lvl, "code": "200 Level"}).json()
        xlsx = c.get(f"/api/templates/{course['template_id']}/export", params={"format": "xlsx"})
        assert xlsx.status_code == 200
        sheet = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
        rows = {str(r[2]).strip(): r[1] for r in sheet.iter_rows(values_only=True) if r and r[2]}
        assert rows["OKEBEGASO Tizaluziz Vovnuzf"] == "DE"
        assert rows["ONE Person"] == "25/PHARM01/001"
