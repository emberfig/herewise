"""A course's groups: split into A to D, a class for one group, and each student counted against
only the classes for everybody and their own group's.

Asked for on 2026-09-29: "create sub groups ... a,b,c,d for 200 level for a particular course ... and
still be able to use the full roster for some courses".
"""
import time
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from conftest import ROSTER
from fake_zk import FakeRecord
from rollcall import db, groups, live
from rollcall.app import create_app


@pytest.fixture(autouse=True)
def fresh_terminal_memory():
    live._seen.clear()
    yield
    live._seen.clear()


@pytest.fixture
def client(tmp_path, k40_factory):
    app = create_app(tmp_path / "groups.db", k40_factory=k40_factory)
    with TestClient(app) as c:
        c.db_path = tmp_path / "groups.db"
        yield c


def a_course(client, code="PCT 201"):
    client.post("/api/setup", json={"username": "lecturer", "display_name": "Dr Lecturer",
                                    "password": "secret123", "academic_session": "2026/2027"})
    level = next(l["id"] for l in client.get("/api/levels").json() if l["level_number"] == 200)
    if not client.get("/api/courses").json():
        preview = client.post(f"/api/levels/{level}/roster/preview",
                              files={"file": ("roster.csv", ROSTER.read_bytes(), "text/csv")})
        client.post(f"/api/roster-imports/{preview.json()['import_id']}/commit")
    course = client.post("/api/courses", json={"level_id": level, "code": code}).json()
    return course, client.get(f"/api/courses/{course['id']}/students").json()


def split(client, course, count=4):
    r = client.post(f"/api/courses/{course['id']}/groups/split", json={"count": count, "order": "matric"})
    assert r.status_code == 200, r.text
    return r.json()["groups"]


def take(client, fake, course, scanners, group_id=None, week=None):
    body = {"group_id": group_id if group_id is not None else "everybody"}
    if week:
        body["week_number"] = week
    r = client.post(f"/api/courses/{course['id']}/live", json=body)
    assert r.status_code == 200, r.text
    class_id = r.json()["register"]["session"]["id"]
    for s in scanners:
        fake.records.append(FakeRecord(user_id=str(s["device_user_id"]),
                                       timestamp=datetime.now().replace(microsecond=0)))
    live._seen.clear()
    ended = client.post(f"/api/live/{class_id}/end", json={})
    assert ended.status_code == 200, ended.text
    return ended.json()


def members_of(client, course, group):
    data = client.get(f"/api/courses/{course['id']}/groups").json()
    return [s for s in data["students"] if s["group_id"] == group["id"]]


def summary_row(client, course, person_id):
    rows = client.get(f"/api/courses/{course['id']}/summary").json()["students"]
    return next(r for r in rows if r["person_id"] == person_id)


def test_a_split_shares_everyone_out_in_unbroken_runs_that_differ_by_one(client):
    course, students = a_course(client)
    made = split(client, course, 4)
    assert [g["name"] for g in made] == ["Group A", "Group B", "Group C", "Group D"]
    sizes = [g["students"] for g in made]
    assert sum(sizes) == len(students) and max(sizes) - min(sizes) <= 1
    data = client.get(f"/api/courses/{course['id']}/groups").json()
    assert data["ungrouped"] == 0
    # Matric order: every Group A number sorts before every Group B number.
    a = [s["matric_no"] for s in data["students"] if s["group_id"] == made[0]["id"] and s["matric_no"]]
    b = [s["matric_no"] for s in data["students"] if s["group_id"] == made[1]["id"] and s["matric_no"]]
    assert max(map(groups._natural, a)) < min(map(groups._natural, b))


def test_another_course_on_the_same_level_keeps_its_whole_roster(client, fake):
    course, students = a_course(client)
    split(client, course)
    other, _ = a_course(client, "PCT 203")
    assert client.get(f"/api/courses/{other['id']}").json()["groups"] == []
    ended = take(client, fake, other, students[:3])
    assert ended["totals"]["students"] == len(students)


def test_a_group_class_registers_only_that_group_and_nobody_else_is_absent(client, fake):
    course, students = a_course(client)
    a, b = split(client, course, 2)
    in_a, in_b = members_of(client, course, a), members_of(client, course, b)
    ended = take(client, fake, course, in_a[:2] + in_b[:1], group_id=a["id"])
    assert ended["totals"]["students"] == len(in_a)
    assert ended["totals"]["present"] == 2
    # The Group B student who scanned is listed as another group's, not counted.
    assert [o["person_id"] for o in ended["others"]] == [in_b[0]["person_id"]]
    row_b = summary_row(client, course, in_b[1]["person_id"])
    assert row_b["held"] == 0 and row_b["missed"] == 0, "Group B was marked absent from Group A's class"
    row_a = summary_row(client, course, in_a[5]["person_id"])
    assert row_a["held"] == 1 and row_a["missed"] == 1


def test_a_class_for_everybody_counts_for_every_group(client, fake):
    course, students = a_course(client)
    a, b = split(client, course, 2)
    take(client, fake, course, students[:1])
    take(client, fake, course, members_of(client, course, a)[:1], group_id=a["id"])
    in_a, in_b = members_of(client, course, a), members_of(client, course, b)
    assert summary_row(client, course, in_a[3]["person_id"])["held"] == 2
    assert summary_row(client, course, in_b[3]["person_id"])["held"] == 1
    summary = client.get(f"/api/courses/{course['id']}/summary").json()
    assert summary["sessions_held"] == 2, "the course figure is the most classes one student had"


def test_the_course_counts_every_class_taken_and_each_student_their_own_classes_left(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    plan = client.put(f"/api/courses/{course['id']}/attendance-plan",
                      json={"weeks": 10, "classes_per_week": 1, "threshold_percent": 75})
    assert plan.status_code == 200, plan.text
    take(client, fake, course, [])
    take(client, fake, course, [], group_id=a["id"])
    take(client, fake, course, [], group_id=a["id"])
    figures = client.get(f"/api/courses/{course['id']}?with_figures=true").json()["figures"]
    # The Classes tab lists three; the most one student was expected at is also three (Group A's).
    assert figures["classes_taken"] == 3 and figures["classes_held"] == 3
    in_b = members_of(client, course, b)[0]
    row_b = summary_row(client, course, in_b["person_id"])
    assert row_b["held"] == 1 and row_b["classes_remaining"] == 9, "Group B's student was given Group A's classes"
    row_a = summary_row(client, course, members_of(client, course, a)[0]["person_id"])
    assert row_a["classes_remaining"] == 7


def test_a_course_with_groups_will_not_start_a_class_without_being_told_whose(client, fake):
    course, _ = a_course(client)
    split(client, course, 2)
    r = client.post(f"/api/courses/{course['id']}/live", json={})
    assert r.status_code == 400
    assert "which group" in r.json()["error"]


def test_moving_from_today_keeps_the_classes_had_with_the_old_group(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    mover = members_of(client, course, a)[0]
    take(client, fake, course, [mover], group_id=a["id"])
    assert groups.classes_with_group(db.connect(client.db_path), course["id"], mover["person_id"]) == 1

    tomorrow = (datetime.now() + timedelta(days=1)).date().isoformat()
    conn = db.connect(client.db_path)
    groups.move(conn, course["id"], mover["person_id"], b["id"], "from_today", None, day=tomorrow)
    row = summary_row(client, course, mover["person_id"])
    assert row["held"] == 1 and row["attended"] == 1, "the Group A class they came to was lost"
    assert row["group"] == "Group A", "they are in Group B only from tomorrow"


def test_the_wrong_group_moves_the_whole_stretch(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    mover = members_of(client, course, a)[0]
    take(client, fake, course, [], group_id=b["id"])
    r = client.post(f"/api/courses/{course['id']}/groups/move",
                    json={"person_id": mover["person_id"], "group_id": b["id"], "how": "wrong_group"})
    assert r.status_code == 200, r.text
    row = summary_row(client, course, mover["person_id"])
    assert row["group"] == "Group B" and row["held"] == 1 and row["missed"] == 1


def test_a_group_with_a_class_cannot_be_removed_and_neither_can_the_split(client, fake):
    course, _ = a_course(client)
    a, _b = split(client, course, 2)
    take(client, fake, course, [], group_id=a["id"])
    assert client.delete(f"/api/groups/{a['id']}").status_code == 400
    assert client.delete(f"/api/courses/{course['id']}/groups").status_code == 400


def test_groups_with_no_class_can_all_go_and_the_course_splits_again(client):
    course, _ = a_course(client)
    split(client, course, 2)
    assert client.delete(f"/api/courses/{course['id']}/groups").status_code == 200
    assert len(split(client, course, 3)) == 3


def test_a_correction_on_another_groups_class_is_refused(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    ended = take(client, fake, course, [], group_id=a["id"])
    outsider = members_of(client, course, b)[0]
    r = client.post(f"/api/sessions/{ended['session']['id']}/marks",
                    json={"person_id": outsider["person_id"], "status": "present", "note": "came along"})
    assert r.status_code == 400
    assert "not in Group A" in r.json()["error"]


def test_a_class_keeps_its_group_while_a_correction_in_it_would_be_left_on_nobodys_register(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    ended = take(client, fake, course, [])
    in_b = members_of(client, course, b)[0]
    r = client.post(f"/api/sessions/{ended['session']['id']}/marks",
                    json={"person_id": in_b["person_id"], "status": "present", "note": "signed the paper sheet"})
    assert r.status_code == 200, r.text
    refused = client.put(f"/api/sessions/{ended['session']['id']}/group", json={"group_id": a["id"]})
    assert refused.status_code == 400 and "corrected" in refused.json()["error"]
    moved = client.put(f"/api/sessions/{ended['session']['id']}/group", json={"group_id": b["id"]})
    assert moved.status_code == 200 and moved.json()["group"] == "Group B"


def test_a_students_evidence_shows_their_own_classes_and_another_groups_only_where_they_scanned(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    visitor, stayer = members_of(client, course, b)[:2]
    take(client, fake, course, [visitor], group_id=a["id"])
    # A live class's window is to the second, so a class opened in the same second as the last one
    # ended would hold that one's scan too. A second apart, as any two real classes are.
    time.sleep(1.1)
    take(client, fake, course, [], group_id=a["id"])
    take(client, fake, course, [], group_id=b["id"])

    def timeline(person):
        record = client.get(f"/api/students/{person['person_id']}/evidence").json()
        return next(c for c in record["courses"] if c["course_id"] == course["id"])["classes"]

    seen = timeline(visitor)
    assert [c["status"] for c in seen if c["status"] == "other_group"] == ["other_group"]
    assert next(c for c in seen if c["status"] == "other_group")["first_scan"]
    assert len(seen) == 2, "Group A's second practical, which they did not come to, was on their timeline"
    assert len(timeline(stayer)) == 1


def test_each_group_has_its_own_sheet_and_every_group_exports_in_one_file(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    take(client, fake, course, members_of(client, course, a)[:3], group_id=a["id"])
    take(client, fake, course, members_of(client, course, b)[:2], group_id=b["id"])
    template_id = client.get(f"/api/courses/{course['id']}").json()["template_id"]
    sheet_a = client.get(f"/api/templates/{template_id}/sheet?group={a['id']}").json()
    assert sheet_a["group"]["name"] == "Group A"
    assert len(sheet_a["students"]) == a["students"]
    held = [s for w in sheet_a["weeks"] for s in w["slots"] if s["state"] == "held"]
    assert len(held) == 1 and held[0]["present"] == 3, "Group A's sheet held Group B's class"
    # One box a week is enough: each group had one class in it.
    assert sheet_a["weeks"][0]["classes"] == 1
    pdf = client.get(f"/api/templates/{template_id}/export?format=pdf&group=all")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    xlsx = client.get(f"/api/templates/{template_id}/export?format=xlsx&group=all")
    assert xlsx.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook
    assert load_workbook(BytesIO(xlsx.content)).sheetnames == ["Group A", "Group B"]


def test_every_group_exports_when_the_students_in_no_group_are_all_carry_over(client, fake):
    # The in-no-group sheet's first row is then the merged CARRY-OVER STUDENTS line, and the
    # spreadsheet's freeze panes failed on it with a 500 (check_exports.mjs on the built app).
    course, _ = a_course(client)
    split(client, course, 2)
    client.post(f"/api/courses/{course['id']}/students",
                json={"entries": [{"matric_no": "PHA/22/9999", "name": "LATE Joiner", "level": "300"}]})
    template_id = client.get(f"/api/courses/{course['id']}").json()["template_id"]
    xlsx = client.get(f"/api/templates/{template_id}/export?format=xlsx&group=all")
    assert xlsx.status_code == 200, xlsx.text[:200]
    from io import BytesIO

    from openpyxl import load_workbook
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames == ["Group A", "Group B", "Students In No Group"]
    assert book["Students In No Group"].freeze_panes, "the names and matric numbers no longer stay in view"


def test_a_class_found_in_the_scans_goes_to_the_group_that_scanned(client, fake):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    base = datetime.now().replace(microsecond=0) - timedelta(days=1)
    for n, s in enumerate(members_of(client, course, b)[:8]):
        fake.records.append(FakeRecord(user_id=str(s["device_user_id"]), timestamp=base + timedelta(minutes=n)))
    pulled = client.post("/api/devices/1/pull")
    assert pulled.status_code == 200, pulled.text
    classes = client.get(f"/api/sessions?course_id={course['id']}").json()
    assert [c["group_id"] for c in classes] == [b["id"]]


def test_pasting_a_department_list_moves_those_students(client):
    course, _ = a_course(client)
    a, b = split(client, course, 2)
    wanted = members_of(client, course, a)[:2]
    lines = [f"{n}, {s['matric_no']}, {s['surname']}" for n, s in enumerate(wanted, start=1)] + ["PHA/99/0000"]
    plan = client.post(f"/api/groups/{b['id']}/list/preview", json={"lines": lines}).json()
    assert plan["move"] == 2 and plan["not_found"] == 1
    client.post(f"/api/groups/{b['id']}/list", json={"lines": lines})
    now_in_b = {s["person_id"] for s in members_of(client, course, b)}
    assert {s["person_id"] for s in wanted} <= now_in_b


def test_a_student_added_after_the_split_is_in_no_group_until_shared_out(client):
    course, _ = a_course(client)
    split(client, course, 2)
    client.post(f"/api/courses/{course['id']}/students",
                json={"entries": [{"matric_no": "PHA/22/9999", "name": "LATE Joiner", "level": "300"}]})
    assert client.get(f"/api/courses/{course['id']}/groups").json()["ungrouped"] == 1
    assert client.post(f"/api/courses/{course['id']}/groups/share-out").json()["placed"] == 1
    assert client.get(f"/api/courses/{course['id']}/groups").json()["ungrouped"] == 0
