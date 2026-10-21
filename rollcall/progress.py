"""Where every student stands part way through the course, and who is at risk.

Two lists come out of the same figures:

- **the notice**, matric numbers and percentages only, to post on a board from any week. Universities
  publish attendance lists before examinations and refuse appeals below the pass mark (the University
  of Uyo's July 2026 directive is the example this follows), so a student has to be able to see where
  they stand while there is still time to attend. Names are left off: it is posted in public.
- **the at-risk list**, for the lecturer, with names, what each student still needs, and the best
  percentage they can still reach if they attend every remaining class. Someone who can no longer
  reach the pass mark should hear it early, not at the examination hall door.

The arithmetic is `register.course_summary`, which already works out what each student needs; this
module states it in words and prints it. Decided with the owner on 2026-09-20: students may see their
percentages during the semester.
"""
from __future__ import annotations

import io
import re
import sqlite3

from . import db, names, pdf_fonts, register

from . import NAME

STATUS_TEXT = {
    "met": "MET",
    "on_track": "ON TRACK",
    "at_risk": "AT RISK",
    "cannot_reach": "CANNOT REACH",
    "below": "BELOW",
    "no_classes_yet": "NO CLASSES YET",
    "all_excused": "ALL EXCUSED",
}
AT_RISK_STATUSES = ("at_risk", "cannot_reach", "below")
VARIANTS = ("notice", "at_risk")


def progress(conn: sqlite3.Connection, course_id: int, date_from: str | None = None,
             date_to: str | None = None) -> dict:
    summary = register.course_summary(conn, course_id, date_from, date_to)
    threshold = summary["threshold_percent"]
    rows = []
    for sn, s in enumerate(summary["students"], start=1):
        rows.append({
            "sn": sn, "person_id": s["person_id"], "device_user_id": s["device_user_id"],
            "matric_no": s["matric_no"], "full_name": f"{s['surname']} {s['other_names']}".strip(),
            "level": s["level"], "carry_over": s["kind"] == "carry_over",
            "attended": s["attended"], "excused": s["excused"], "counted": s["counted"],
            "percent": s["percent"], "status": s["status"], "status_text": STATUS_TEXT[s["status"]],
            "more_needed": s["more_needed"], "can_still_miss": s["can_still_miss"],
            "best_possible_percent": s["best_possible_percent"],
        })
    counts = {key: sum(1 for r in rows if r["status"] == key) for key in STATUS_TEXT}
    excluded = summary["plan"]["excused_rule"] != "present"
    return {
        "course": summary["course"],
        "rules": {
            "threshold_percent": threshold,
            "excused_excluded": excluded,
            "classes_held": summary["sessions_held"],
            "planned_classes": summary["planned_classes"],
            "classes_remaining": summary["classes_remaining"],
            "date_from": date_from,
            "date_to": date_to,
        },
        "counts": {**counts, "total": len(rows),
                   "at_risk_list": sum(1 for r in rows if r["status"] in AT_RISK_STATUSES)},
        "class_average_percent": summary["class_average_percent"],
        "warnings": summary["warnings"],
        "students": rows,
        "generated_at": db.now(),
    }


def rows_for(data: dict, variant: str) -> list[dict]:
    if variant == "at_risk":
        return [r for r in data["students"] if r["status"] in AT_RISK_STATUSES]
    return data["students"]


def _heading(data: dict, variant: str) -> list[str]:
    c, r = data["course"], data["rules"]
    remaining = r["classes_remaining"]
    period = ""
    if r["date_from"] or r["date_to"]:
        period = f", {r['date_from'] or 'start'} to {r['date_to'] or 'today'}"
    whose = "your" if variant == "notice" else "the student's"
    excused = (f"Excused absences are left out of {whose} total" if r["excused_excluded"]
               else "Excused absences count as attended")
    lines = [
        f"COURSE: {c['code']}{(' ' + c['title']) if c['title'] else ''}",
        f"{c['level'].upper()}, {c['academic_session']} ACADEMIC SESSION",
        f"Pass mark {r['threshold_percent']}% of classes. {r['classes_held']} classes held so far{period}"
        + (f", {remaining} still to come." if remaining else ".") + f" {excused}.",
    ]
    if variant == "notice":
        lines.append("This is where you stand today, not a final list. It changes as more classes are held.")
    return lines


def _columns(variant: str) -> list[tuple[str, float, float]]:
    """(label, PDF width in points, Excel column width)."""
    if variant == "notice":
        return [("S/N", 34, 6), ("MATRIC NO.", 150, 20), ("ATTENDED", 70, 12), ("OF", 50, 8),
                ("%", 60, 8), ("STATUS", 110, 16)]
    return [("S/N", 28, 6), ("MATRIC NO.", 96, 18), ("FULL NAME", 150, 32), ("LEVEL", 44, 10),
            ("ATTENDED", 52, 12), ("OF", 34, 8), ("%", 40, 8), ("BEST POSSIBLE", 66, 15),
            ("STILL NEEDS", 58, 13), ("STATUS", 76, 16)]


def _values(row: dict, variant: str) -> list[str]:
    percent = f"{row['percent']:g}"
    if variant == "notice":
        return [str(row["sn"]), names.matric_text(row["matric_no"]), str(row["attended"]), str(row["counted"]),
                percent, row["status_text"]]
    level = (f"{row['level']}L" if row["level"] else "") + (" CO" if row["carry_over"] else "")
    best = "" if row["best_possible_percent"] is None else f"{row['best_possible_percent']:g}"
    needs = "" if row["more_needed"] is None else str(row["more_needed"])
    return [str(row["sn"]), names.matric_text(row["matric_no"]), row["full_name"], level.strip(), str(row["attended"]),
            str(row["counted"]), percent, best, needs, row["status_text"]]


def _title(data: dict, variant: str) -> str:
    if variant == "notice":
        return "ATTENDANCE SO FAR"
    return f"STUDENTS BELOW {data['rules']['threshold_percent']}% SO FAR"


def _summary_line(data: dict, variant: str) -> str:
    c = data["counts"]
    if variant == "notice":
        return (f"{c['total']} students. On track or met: {c['met'] + c['on_track']}. "
                f"At risk: {c['at_risk']}. Cannot reach the pass mark: {c['cannot_reach']}.")
    return (f"{c['at_risk_list']} of {c['total']} students are below the pass mark. "
            f"Cannot reach it any more: {c['cannot_reach']}. Class average {data['class_average_percent']:g}%.")


def to_pdf(data: dict, variant: str = "notice", paper: str = "a4") -> bytes:
    from reportlab.lib.colors import Color, black
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    regular, bold = pdf_fonts.families()
    page_w, page_h = pdf_fonts.page_size(paper)
    left, right = 36.0, page_w - 36.0
    row_h = 14.0
    # Scaled to the page: the at-risk list's ten columns are wider than A4 at their natural widths,
    # which pushed S/N off the left edge and cut STATUS off the right (caught by rendering it).
    natural = [(label, width) for label, width, _ in _columns(variant)]
    scale = min(1.0, (right - left) / sum(w for _, w in natural))
    cols = [(label, width * scale) for label, width in natural]
    table_w = sum(w for _, w in cols)
    x0 = left + (right - left - table_w) / 2
    pink = Color(0.97, 0.84, 0.85)
    rows = rows_for(data, variant)

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(page_w, page_h))
    pdf.setTitle(f"{_title(data, variant)} {data['course']['code']}")
    pdf.setAuthor(NAME)

    def fit(value: str, x: float, y: float, width: float, font: str, size: float, align: str) -> None:
        value = pdf_fonts.text(value)
        while size > 5 and stringWidth(value, font, size) > width - 4:
            size -= 0.3
        pdf.setFont(font, size)
        if align == "centre":
            pdf.drawCentredString(x + width / 2, y, value)
        else:
            pdf.drawString(x + 3, y, value)

    def header(page: int) -> float:
        pdf.setFillColor(black)
        pdf.setFont(bold, 13)
        pdf.drawCentredString(page_w / 2, page_h - 48, _title(data, variant))
        y = page_h - 64
        for line in _heading(data, variant):
            fit(line, left, y, right - left, bold if line.startswith("COURSE") else regular, 9, "centre")
            y -= 13
        pdf.setFont(regular, 8)
        pdf.drawRightString(right, 24, f"Page {page}")
        pdf.drawString(left, 24, f"Generated by {NAME} on {data['generated_at'].replace('T', ' ')}")
        y -= 6
        x = x0
        pdf.setLineWidth(0.8)
        for label, width in cols:
            pdf.rect(x, y - row_h, width, row_h)
            fit(label, x, y - row_h + 4, width, bold, 7.5, "centre")
            x += width
        return y - row_h

    page = 1
    y = header(page)
    for i, row in enumerate(rows):
        last_row = i == len(rows) - 1
        if y - row_h < 40 or (last_row and y - row_h < 60):
            pdf.showPage()
            page += 1
            y = header(page)
        x = x0
        pdf.setLineWidth(0.4)
        for (label, width), value in zip(cols, _values(row, variant)):
            if label == "STATUS" and row["status"] == "cannot_reach":
                pdf.setFillColor(pink)
                pdf.rect(x, y - row_h, width, row_h, fill=1, stroke=0)
                pdf.setFillColor(black)
            pdf.rect(x, y - row_h, width, row_h)
            emphasis = label == "STATUS" and row["status"] in AT_RISK_STATUSES
            fit(value, x, y - row_h + 4, width, bold if emphasis else regular, 7.5,
                "left" if label == "FULL NAME" else "centre")
            x += width
        y -= row_h

    if not rows:
        fit("Nobody is below the pass mark." if variant == "at_risk" else "Nobody on this course yet.",
            x0, y - 20, table_w, regular, 9, "centre")
        y -= 20
    y -= 16
    pdf.setFont(bold, 9)
    pdf.drawString(x0, y, _summary_line(data, variant))
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def to_xlsx(data: dict, variant: str = "notice") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance so far" if variant == "notice" else "At risk"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(name="Arial", bold=True, size=9)
    normal = Font(name="Arial", size=9)
    centre = Alignment(horizontal="center", vertical="center")
    pink = PatternFill("solid", fgColor="F8D7DA")

    cols = _columns(variant)
    last = len(cols)
    lines = [_title(data, variant)] + _heading(data, variant)
    for i, text in enumerate(lines, start=1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=last)
        cell = ws.cell(i, 1, text)
        cell.font = Font(name="Arial", bold=True, size=12 if i == 1 else 9)
        cell.alignment = centre
    header_row = len(lines) + 2
    for col, (label, _, width) in enumerate(cols, start=1):
        cell = ws.cell(header_row, col, label)
        cell.font, cell.alignment, cell.border = bold, centre, box
        ws.column_dimensions[cell.column_letter].width = width

    r = header_row
    for row in rows_for(data, variant):
        r += 1
        for col, ((label, _, _), value) in enumerate(zip(cols, _values(row, variant)), start=1):
            cell = ws.cell(r, col, int(value) if value.isdigit() else value)
            cell.font = bold if (label == "STATUS" and row["status"] in AT_RISK_STATUSES) else normal
            cell.border, cell.alignment = box, centre
            if label == "STATUS" and row["status"] == "cannot_reach":
                cell.fill = pink
    ws.cell(r + 2, 1, _summary_line(data, variant)).font = bold

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def filename(data: dict, extension: str, variant: str) -> str:
    code = re.sub(r"[^A-Za-z0-9]+", "-", data["course"]["code"] or "course").strip("-")
    kind = "attendance-so-far" if variant == "notice" else "at-risk"
    return f"{code}-{kind}-{data['generated_at'][:10]}.{extension}"
