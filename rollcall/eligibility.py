"""The examination eligibility list: the decision the lecturer is accountable for.

Every row says how the number was reached (attended, excused, classes counted), and the list itself
states the rules used: the pass mark, how excused absences were treated, how many classes were held
and the date range. A list produced before the planned classes are all held is marked PROVISIONAL.

Decisions:
  eligible      percentage at or above the pass mark
  not_eligible  below it
  review        nothing to judge: no classes held, or every class excused
"""
from __future__ import annotations

import sqlite3

from . import db, register

from . import NAME, names

VARIANTS = ("full", "notice")


def eligibility(conn: sqlite3.Connection, course_id: int, date_from: str | None = None,
                date_to: str | None = None) -> dict:
    summary = register.course_summary(conn, course_id, date_from, date_to)
    threshold = summary["threshold_percent"]
    rules_excused = summary["plan"]["excused_rule"]
    rows = []
    for sn, s in enumerate(summary["students"], start=1):
        if s["status"] in ("no_classes_yet", "all_excused") or not s["held"]:
            decision = "review"
        else:
            decision = "eligible" if s["percent"] >= threshold else "not_eligible"
        rows.append({"sn": sn, "person_id": s["person_id"], "device_user_id": s["device_user_id"],
                     "matric_no": s["matric_no"], "full_name": f"{s['surname']} {s['other_names']}".strip(),
                     "level": s["level"], "carry_over": s["kind"] == "carry_over", "group": s.get("group"),
                     "attended": s["attended"],
                     "excused": s["excused"], "counted": s["counted"], "percent": s["percent"],
                     "decision": decision})

    provisional = bool(summary["classes_remaining"])
    excused_text = ("Excused absences are left out of the student's total" if rules_excused != "present"
                    else "Excused absences count as attended")
    return {
        "course": summary["course"],
        "rules": {
            "threshold_percent": threshold,
            "excused_rule": rules_excused,
            "excused_text": excused_text,
            "classes_held": summary["sessions_held"],
            "grouped": summary.get("grouped", False),
            "planned_classes": summary["planned_classes"],
            "classes_remaining": summary["classes_remaining"],
            "date_from": date_from,
            "date_to": date_to,
            "provisional": provisional,
        },
        "counts": {
            "total": len(rows),
            "eligible": sum(1 for r in rows if r["decision"] == "eligible"),
            "not_eligible": sum(1 for r in rows if r["decision"] == "not_eligible"),
            "review": sum(1 for r in rows if r["decision"] == "review"),
            "carry_over": sum(1 for r in rows if r["carry_over"]),
        },
        "warnings": summary["warnings"],
        "students": rows,
        "generated_at": db.now(),
    }


DECISION_TEXT = {"eligible": "ELIGIBLE", "not_eligible": "NOT ELIGIBLE", "review": "REVIEW"}


def _heading(data: dict) -> list[str]:
    c, r = data["course"], data["rules"]
    period = ""
    if r["date_from"] or r["date_to"]:
        period = f", {r['date_from'] or 'start'} to {r['date_to'] or 'today'}"
    held = (f"Up to {r['classes_held']} classes held for each student{period}, counting the classes for "
            "everybody and their own group's only" if r.get("grouped")
            else f"{r['classes_held']} classes held{period}")
    return [
        f"COURSE: {c['code']}{(' ' + c['title']) if c['title'] else ''}",
        f"{c['level'].upper()}, {c['academic_session']} ACADEMIC SESSION",
        f"Pass mark {r['threshold_percent']}% of classes. {held}. {r['excused_text']}.",
    ]


def to_xlsx(data: dict, variant: str = "full") -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Eligibility"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(name="Arial", bold=True, size=9)
    normal = Font(name="Arial", size=9)
    centre = Alignment(horizontal="center", vertical="center")
    red = PatternFill("solid", fgColor="F8D7DA")

    if variant == "notice":
        headers = ["S/N", "MATRIC NO.", "STATUS"]
        widths = [6, 20, 16]
    else:
        headers = ["S/N", "MATRIC NO.", "FULL NAME", "LEVEL", "ATTENDED", "EXCUSED", "COUNTED", "%", "STATUS"]
        widths = [6, 18, 36, 12, 10, 10, 10, 8, 15]
    last = len(headers)

    title = "EXAMINATION ELIGIBILITY LIST" + (" (PROVISIONAL)" if data["rules"]["provisional"] else "")
    lines = [title] + _heading(data)
    for i, text in enumerate(lines, start=1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=last)
        cell = ws.cell(i, 1, text)
        cell.font = Font(name="Arial", bold=True, size=12 if i == 1 else 9)
        cell.alignment = centre
    header_row = len(lines) + 2
    for col, (text, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(header_row, col, text)
        cell.font, cell.alignment, cell.border = bold, centre, box
        ws.column_dimensions[cell.column_letter].width = width

    r = header_row
    for s in data["students"]:
        r += 1
        if variant == "notice":
            values = [s["sn"], names.matric_text(s["matric_no"]), DECISION_TEXT[s["decision"]]]
        else:
            level = f"{s['level']}L" + (" CO" if s["carry_over"] else "") if s["level"] else ("CO" if s["carry_over"] else "")
            values = [s["sn"], names.matric_text(s["matric_no"]), s["full_name"], level, s["attended"], s["excused"], s["counted"],
                      s["percent"], DECISION_TEXT[s["decision"]]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            cell.font, cell.border = normal, box
            cell.alignment = Alignment(horizontal="left" if headers[col - 1] == "FULL NAME" else "center")
            if s["decision"] == "not_eligible" and headers[col - 1] == "STATUS":
                cell.fill = red
                cell.font = bold

    counts = data["counts"]
    r += 2
    ws.cell(r, 1, f"Total {counts['total']}, eligible {counts['eligible']}, not eligible {counts['not_eligible']}, "
                  f"review {counts['review']}, carry-over {counts['carry_over']}.").font = bold
    if variant == "full":
        for label in ("COURSE LECTURER", "HEAD OF DEPARTMENT"):
            r += 3
            ws.cell(r, 1, f"{label}:  Name ______________________   Signature ______________   Date __________").font = bold
    r += 2
    ws.cell(r, 1, f"Generated by {NAME} on {data['generated_at'].replace('T', ' ')}").font = Font(name="Arial", size=8)
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(data: dict, variant: str = "full", paper: str = "a4") -> bytes:
    import io

    from reportlab.lib.colors import Color, black
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    from .pdf_fonts import families, page_size, text

    regular, bold = families()
    page_w, page_h = page_size(paper)
    left, right = 36.0, page_w - 36.0
    row_h = 14.0
    if variant == "notice":
        cols = [("S/N", 40), ("MATRIC NO.", 160), ("STATUS", 120)]
    else:
        cols = [("S/N", 26), ("MATRIC NO.", 84), ("FULL NAME", 170), ("LEVEL", 44), ("ATT.", 36), ("EXC.", 36),
                ("OF", 34), ("%", 36), ("STATUS", 57)]
    table_w = sum(w for _, w in cols)
    x0 = left + (right - left - table_w) / 2
    pink = Color(0.97, 0.84, 0.85)

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=(page_w, page_h))
    pdf.setTitle(f"Eligibility {data['course']['code']}")
    pdf.setAuthor(NAME)

    def fit(value: str, x: float, y: float, width: float, font: str, size: float, align: str) -> None:
        value = text(value)
        while size > 5 and stringWidth(value, font, size) > width - 4:
            size -= 0.3
        pdf.setFont(font, size)
        if align == "centre":
            pdf.drawCentredString(x + width / 2, y, value)
        else:
            pdf.drawString(x + 3, y, value)

    def header(page: int) -> float:
        title = "EXAMINATION ELIGIBILITY LIST" + (" (PROVISIONAL)" if data["rules"]["provisional"] else "")
        pdf.setFillColor(black)
        pdf.setFont(bold, 13)
        pdf.drawCentredString(page_w / 2, page_h - 48, title)
        y = page_h - 64
        for line in _heading(data):
            fit(line, left, y, right - left, bold if line.startswith("COURSE") else regular, 9, "centre")
            y -= 13
        pdf.setFont(regular, 8)
        pdf.drawRightString(right, 24, f"Page {page}")
        pdf.drawString(left, 24, f"Generated by {NAME} on {data['generated_at'].replace('T', ' ')}")
        y -= 6
        x = x0
        pdf.setLineWidth(0.8)
        for label, width in cols:
            pdf.rect(x, y - row_h, width, row_h)
            fit(label, x, y - row_h + 4, width, bold, 7.5, "centre")
            x += width
        return y - row_h

    page = 1
    y = header(page)
    footer_space = 120 if variant == "full" else 50
    students = data["students"]
    for i, s in enumerate(students):
        last_row = i == len(students) - 1
        if y - row_h < 40 or (last_row and y - row_h < footer_space):
            pdf.showPage()
            page += 1
            y = header(page)
        if variant == "notice":
            values = [str(s["sn"]), names.matric_text(s["matric_no"]), DECISION_TEXT[s["decision"]]]
        else:
            level = (f"{s['level']}L" if s["level"] else "") + (" CO" if s["carry_over"] else "")
            values = [str(s["sn"]), names.matric_text(s["matric_no"]), s["full_name"], level.strip(), str(s["attended"]),
                      str(s["excused"]), str(s["counted"]), f"{s['percent']:g}", DECISION_TEXT[s["decision"]]]
        x = x0
        pdf.setLineWidth(0.4)
        for (label, width), value in zip(cols, values):
            if label == "STATUS" and s["decision"] == "not_eligible":
                pdf.setFillColor(pink)
                pdf.rect(x, y - row_h, width, row_h, fill=1, stroke=0)
                pdf.setFillColor(black)
            pdf.rect(x, y - row_h, width, row_h)
            emphasis = label == "STATUS" and s["decision"] != "eligible"
            fit(value, x, y - row_h + 4, width, bold if emphasis else regular, 7.5,
                "left" if label == "FULL NAME" else "centre")
            x += width
        y -= row_h

    counts = data["counts"]
    y -= 16
    pdf.setFont(bold, 9)
    pdf.drawString(x0, y, f"Total {counts['total']}. Eligible {counts['eligible']}. Not eligible "
                          f"{counts['not_eligible']}. Review {counts['review']}. Carry-over {counts['carry_over']}.")
    if variant == "full":
        for label in ("COURSE LECTURER", "HEAD OF DEPARTMENT"):
            y -= 34
            pdf.setFont(bold, 8.5)
            pdf.drawString(x0, y, label)
            pdf.setFont(regular, 8.5)
            pdf.drawString(x0 + 120, y, "Name ____________________   Signature ______________   Date __________")
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def filename(data: dict, extension: str, variant: str) -> str:
    import re

    base = f"Eligibility {data['course']['code']} {data['course']['academic_session']}"
    if variant == "notice":
        base += " notice"
    return re.sub(r"[^A-Za-z0-9 ._-]+", "-", base).strip() + f".{extension}"
