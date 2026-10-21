"""Roster import (PRD section 7). Always two steps: a preview that changes nothing, then a commit.

A student is identified by matric number. Re-importing the same student keeps their device ID, so a
roster for next session moves them up a level without touching the device or their fingerprints.
IDs are allocated from the level's block and are never reused, because old punches refer to them.
"""
from __future__ import annotations

import csv
import io
import re
import sqlite3
import zipfile
from collections import Counter
from collections.abc import Iterable
from dataclasses import asdict, dataclass, field

from . import courses, db, levels, names

REQUIRED_COLUMNS = ("matric_no", "name")
DEFAULT_USER_CAPACITY = 1000

# A roster is a few kilobytes: the department's largest list, 251 students, is 16 KB as .xlsx. Five
# megabytes leaves room for any real one and turns away a video or a backup chosen by mistake before
# it is read. A 20 MB file of text used to stop the reader with an error nobody could act on
# (2026-09-24).
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

# What the university's own exports call these columns. The header is reduced to lower-case words
# before it is looked up here, so "Matric No.", "MATRIC NO", "Matric  Number" and "matric_no" all
# arrive as the same key. The university's course list headers its second column "Matric No.", and the full
# stop alone was enough to have the file refused with "Missing column(s): matric_no" (2026-09-21).
ALIASES = {
    "matric no": "matric_no",
    "matric number": "matric_no",
    "matric": "matric_no",
    "matriculation no": "matric_no",
    "matriculation number": "matric_no",
    "reg no": "matric_no",
    "registration no": "matric_no",
    "registration number": "matric_no",
    "full name": "name",
    "student name": "name",
    "names": "name",
    "student": "name",
    # A name kept over two or three columns, as many departments keep it. "S/N, MATRIC NO, SURNAME,
    # OTHER NAMES" was refused as "Missing column(s): name" (2026-09-24).
    "last name": "surname",
    "lastname": "surname",
    "family name": "surname",
    "other name": "other_names",
    "othernames": "other_names",
    "given names": "other_names",
    "first names": "other_names",
    "forenames": "other_names",
    "firstname": "first_name",
    "given name": "first_name",
    "forename": "first_name",
    "middle names": "middle_name",
    "middlename": "middle_name",
}

# The parts of a name after the surname, when each has a column of its own, in the order they join.
GIVEN_COLUMNS = ("first_name", "middle_name", "other_names")

# The first eight bytes of an OLE2 compound file, which is what a real .xls is, and of a zip, which
# is what .xlsx is. Read from the bytes rather than the file name, because a name can lie.
XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"

# The first bytes of the files a person picks by mistake, so the refusal can say what was picked.
PICTURES = ((b"\x89PNG", "a picture (PNG)"), (b"\xff\xd8\xff", "a picture (JPEG)"),
            (b"GIF8", "a picture (GIF)"))


@dataclass
class PlannedStudent:
    line: int
    matric_no: str
    source_name: str
    surname: str
    other_names: str
    device_name: str
    shortened: bool
    action: str                         # new | update | unchanged | skip
    device_user_id: int | None = None
    person_id: int | None = None
    issues: list[str] = field(default_factory=list)


class RosterError(ValueError):
    pass


class NeedsColumns(RosterError):
    """The file was read, but which column holds what has to come from the person importing it.

    Raised where a refusal used to be the end of it: headings that name no matric or name column,
    columns that hold the wrong thing, a workbook with several lists and none named for the level.
    The import screen answers it by showing the first rows and asking; anywhere else it is an
    ordinary refusal with the same sentence.
    """

    def __init__(self, reason: str, table: list[list[str]], guess: dict | None = None,
                 sheets: list[str] | None = None, sheet: str | None = None):
        super().__init__(reason)
        self.reason = reason
        self.table = table
        self.guess = guess if guess is not None else guess_columns(table)
        self.sheets = sheets or []
        self.sheet = sheet

    def payload(self) -> dict:
        """What the import screen needs to ask: a few rows around the first student, and a guess."""
        first = self.guess.get("first_row") or 1
        start = max(0, first - 1 - SAMPLE_BEFORE)
        rows = []
        for index in range(start, len(self.table)):
            cells = _trim(self.table[index] or [])
            if not any(c.strip() for c in cells):
                continue
            rows.append({"row": index + 1, "cells": [c.strip()[:SAMPLE_CELL] for c in cells[:SAMPLE_COLUMNS]]})
            if len(rows) == SAMPLE_ROWS:
                break
        width = max((len(r["cells"]) for r in rows), default=0)
        for r in rows:
            r["cells"] += [""] * (width - len(r["cells"]))
        return {"needs_columns": True, "reason": self.reason, "sheets": self.sheets, "sheet": self.sheet,
                "columns": width, "rows": rows, "guess": self.guess}


# How much of a file the question shows: enough rows to recognise the list and its headings, and
# no more columns than a laptop screen holds.
SAMPLE_BEFORE, SAMPLE_ROWS, SAMPLE_COLUMNS, SAMPLE_CELL = 3, 10, 15, 40


@dataclass
class Choice:
    """Which column is which, chosen on the import screen. Columns count from 0, rows from 1, as the
    screen showed them. `other` is the column holding the other names when the name is split."""
    matric: int
    name: int
    first_row: int
    other: int | None = None
    level: int | None = None


def decode(raw: bytes) -> str:
    # Excel's "Unicode Text" is UTF-16 with tabs, and was read as "m a t r i c _ n o" (2026-09-24).
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    head = raw[:400]
    if len(head) >= 8 and head[1::2].count(0) * 2 > len(head[1::2]):
        return raw.decode("utf-16-le", errors="replace")
    if len(head) >= 8 and head[0::2].count(0) * 2 > len(head[0::2]):
        return raw.decode("utf-16-be", errors="replace")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise RosterError("The file is not readable text. Save it as CSV (UTF-8) and try again.")


def column_name(header: str) -> str:
    """One header cell as a column name. "Matric No." and "MATRIC  NUMBER" are both matric_no."""
    words = re.sub(r"[^a-z0-9]+", " ", str(header or "").lower()).strip()
    return ALIASES.get(words, words.replace(" ", "_"))


def _trim(cells: list[str]) -> list[str]:
    """Drop trailing empty cells. A spreadsheet saved as CSV carries its unused columns with it,
    so the university's course list arrives as "S/N,Matric No.,Name,,,,," and every row five commas long."""
    end = len(cells)
    while end and not str(cells[end - 1] or "").strip():
        end -= 1
    return [str(c or "") for c in cells[:end]]


# How far down a file to look for the row that names the columns. A departmental spreadsheet puts
# the university's name, the session and the title above it: the one this was written for had five
# such rows and a blank one before "S/N, NAMES, MATRIC NO., GROUP". Twenty is far more than any of
# them and still refuses a file that simply has no header at all.
HEADER_SEARCH_ROWS = 20


def _names_the_columns(columns: list[str]) -> bool:
    return "matric_no" in columns and ("name" in columns or "surname" in columns)


def find_header(table: list[list[str]]) -> int:
    """Which row names the columns. The first one that names both of the columns we need.

    Demanding the very first row was right for a file straight from the portal and wrong for one a
    department made: every sheet of "Level groupings.xlsx" opened with the university's name and
    the session, and the import refused all five of them (reported 2026-09-22).
    """
    for index, row in enumerate(table[:HEADER_SEARCH_ROWS]):
        if _names_the_columns([column_name(h) for h in _trim(row or [])]):
            return index
    return -1


# -- what a column holds -----------------------------------------------------------------------
#
# Read from the values, not the heading, because the heading is what was wrong. A matric number
# has a digit in it and a slash or a dash ("23/PHARX01/015"), or is a long run of digits. A name
# has letters and no digits.

_MATRIC_SHAPE = re.compile(r"\d.*[/\-]|[/\-].*\d|^\d{5,}$")


def looks_like_matric(value: str) -> bool:
    return bool(_MATRIC_SHAPE.search(value.strip()))


def looks_like_name(value: str) -> bool:
    value = value.strip()
    return sum(ch.isalpha() for ch in value) >= 2 and not any(ch.isdigit() for ch in value)


def column_trouble(rows: list[dict]) -> str | None:
    """Why the matric and name columns do not hold what their headings say, or None.

    Headings the right way round over values the wrong way round imported six students called
    "25/TEST01/001" with no complaint at all (2026-09-24). More than half of a column has to be
    wrong before a file is questioned, so one odd row is left to the row checks.
    """
    matrics = [r["matric_no"] for r in rows
               if r.get("matric_no", "").strip() and not names.is_direct_entry(r["matric_no"])]
    full = [r.get("name", "") for r in rows if r.get("name", "").strip()]
    no_digits = [m for m in matrics if not any(ch.isdigit() for ch in m)]
    named_like_matric = [n for n in full if looks_like_matric(n) or not looks_like_name(n)]
    matric_wrong = matrics and len(no_digits) * 2 > len(matrics)
    name_wrong = full and len(named_like_matric) * 2 > len(full)
    if matric_wrong and name_wrong:
        return (f"The matric number and name columns look the wrong way round: the matric number "
                f"column holds \"{no_digits[0]}\" and the name column holds \"{named_like_matric[0]}\".")
    if matric_wrong:
        return f"The matric number column holds \"{no_digits[0]}\", which has no number in it."
    if name_wrong:
        return f"The name column holds \"{named_like_matric[0]}\", which is not a name."
    return None


def guess_columns(table: list[list[str]], header_at: int = -1) -> dict:
    """A first answer to "which column is which", for the screen to offer and the person to correct.

    Each column is scored on how many of its cells look like a matric number and how many like a
    whole name. The first row is the first whose matric cell looks like one. Nothing here is used
    without the person pressing Preview on it.
    """
    body = [_trim(r or []) for r in table[:200]]
    width = max((len(r) for r in body), default=0)

    def score(test, words=0):
        counts = []
        for c in range(width):
            cells = [r[c] for r in body if c < len(r)]
            counts.append(sum(1 for v in cells if test(v) and (not words or len(v.split()) >= words)))
        return counts

    matric_scores = score(looks_like_matric)
    full_scores = score(looks_like_name, words=2)
    single_scores = score(looks_like_name)
    matric = max(range(width), key=lambda c: matric_scores[c], default=None)
    if matric is not None and not matric_scores[matric]:
        matric = None
    candidates = [c for c in range(width) if c != matric]
    name = max(candidates, key=lambda c: (full_scores[c], single_scores[c]), default=None)
    if name is not None and not single_scores[name]:
        name = None
    other = None
    if name is not None and not full_scores[name]:
        # Only single words in the best column: a surname, with the other names beside it.
        beside = [c for c in candidates if c != name and single_scores[c]]
        other = min(beside, key=lambda c: abs(c - name), default=None)

    first_row = 1
    if header_at >= 0:
        first_row = header_at + 2
    elif matric is not None:
        first_row = next((i + 1 for i, r in enumerate(body)
                          if matric < len(r) and looks_like_matric(r[matric])), 1)
    level = None
    heading = _trim(table[first_row - 2] or []) if first_row >= 2 and first_row - 2 < len(table) else []
    for c, cell in enumerate(heading):
        if column_name(cell) == "level":
            level = c
    return {"matric_col": matric, "name_col": name, "other_col": other, "level_col": level,
            "first_row": first_row}


def _read(table: list[list[str]]) -> tuple[list[dict], int, list[str]]:
    """The rows under the header, the header's position, and the column names it gave."""
    if not any(any(str(f).strip() for f in (row or [])) for row in table):
        raise RosterError("The file is empty.")
    at = find_header(table)
    if at < 0:
        first = _trim(table[0] if table else [])
        named = ", ".join(h.strip() for h in first if h.strip()) or "nothing"
        found = [column_name(h) for h in first]
        missing = [{"matric_no": "matric number", "name": "name"}[c] for c in REQUIRED_COLUMNS
                   if c not in found and not (c == "name" and "surname" in found)]
        raise NeedsColumns(
            f"No row near the top of the file names the {' or the '.join(missing)} column, so "
            f"Herewise cannot tell which column holds what. The first row says: {named}. A matric "
            "number column and a name column are needed; a level column is read if it is there.", table)
    header = _trim(table[at] or [])
    columns = [column_name(h) for h in header]
    if "name" in columns and "surname" not in columns and any(c in columns for c in GIVEN_COLUMNS):
        # NAME beside OTHER NAMES is the surname, whatever it is called.
        columns[columns.index("name")] = "surname"
    split = "name" not in columns
    # An unquoted comma in a name spills it into the next column. It is folded back into the name,
    # or into the last part of it when the name is kept over several columns.
    fold_at = columns.index("name") if not split else max(
        columns.index(c) for c in ("surname", *GIVEN_COLUMNS) if c in columns)
    rows = []
    # Numbered from the header's own position, so "line 61" is row 61 of the file somebody is
    # looking at rather than the 61st row after a header they cannot see.
    for line, raw_fields in enumerate(table[at + 1:], start=at + 2):
        fields = _trim(raw_fields or [])
        if not any(f.strip() for f in fields):
            continue
        surplus = len(fields) - len(columns)
        if surplus > 0:
            joined = ", ".join(f.strip() for f in fields[fold_at:fold_at + surplus + 1])
            fields = fields[:fold_at] + [joined] + fields[fold_at + surplus + 1:]
        elif surplus < 0:
            fields = fields + [""] * -surplus
        row = {c: fields[i].strip() for i, c in enumerate(columns)}
        if split:
            row["other_names"] = " ".join(row[c] for c in GIVEN_COLUMNS if row.get(c))
            row["name"] = " ".join(p for p in (row.get("surname", ""), row["other_names"]) if p)
        row["_line"] = line
        row["_joined"] = surplus > 0
        rows.append(row)
    if not rows:
        named = ", ".join(h.strip() for h in header if h.strip())
        raise RosterError(
            f"This file has the headings ({named}) on row {at + 1} but no students under them. "
            "Check it is the list itself, saved after it was filled in.")
    return rows, at, columns


def parse_rows(table: list[list[str]]) -> list[dict]:
    """Read the header row and the rows under it, repairing names that contain an unquoted comma.

    A row with more fields than the header has its surplus folded back into the name column, so
    "A/1,OZAN-DELE Emmanuel Chibueze, Xuziil,200" still reads the level as 200.
    """
    return _read(table)[0]


def rows_by_choice(table: list[list[str]], choice: Choice) -> list[dict]:
    """The students in a file whose columns the person named, from the first row they named.

    A row with neither a matric number nor a name in the chosen columns is passed over, which is a
    blank line or a total under the list. Everything else is kept for the plan's row checks, so a
    wrong choice shows up as refusals with reasons rather than as nothing.
    """
    rows = []
    for index in range(max(0, choice.first_row - 1), len(table)):
        cells = _trim(table[index] or [])

        def cell(column):
            return cells[column].strip() if column is not None and 0 <= column < len(cells) else ""

        matric, name = cell(choice.matric), cell(choice.name)
        if not matric and not name:
            continue
        row = {"matric_no": matric, "level": cell(choice.level), "_line": index + 1, "_joined": False}
        if choice.other is not None:
            row["surname"], row["other_names"] = name, cell(choice.other)
            row["name"] = " ".join(p for p in (name, row["other_names"]) if p)
        else:
            row["name"] = name
        rows.append(row)
    if not rows:
        raise RosterError(f"Nothing was found from row {choice.first_row} down in the columns chosen.")
    return rows


# Tried in this order, and the first that finds the headings wins. Excel saves "Text (Tab
# delimited)" and "Unicode Text" with tabs, and in some regions a CSV with semicolons; each was
# refused as "Missing column(s)" (2026-09-24). Commas first, because that is what a CSV is.
SEPARATORS = (",", "\t", ";")


def csv_table(text: str) -> list[list[str]]:
    """The rows of a CSV, whatever separates its columns."""
    tables = []
    for separator in SEPARATORS:
        try:
            table = [list(r) for r in csv.reader(io.StringIO(text, newline=""), delimiter=separator)]
        except csv.Error as error:
            # Python's own limit on one cell is 128 KB, which no roster comes near.
            raise RosterError("This file could not be read: one cell in it is longer than a whole "
                              "roster, so it is probably not the list. Check it is the right file.") from error
        if find_header(table) >= 0:
            return table
        tables.append(table)
    # None found the headings. Keep the reading that splits the top rows into the most columns: it is
    # the one to show when asking which column is which.
    return max(tables, key=lambda t: sum(len(_trim(r)) for r in t[:HEADER_SEARCH_ROWS]))


def parse_csv(text: str) -> list[dict]:
    return parse_rows(csv_table(text))


def _cell(value) -> str:
    """A spreadsheet cell as the text a person sees. A level typed as a number comes back as 200.0
    from both readers, and "200.0" is not a level."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


_UNREADABLE = ("This file could not be read as a spreadsheet. Open it in Excel and save it as an "
               "Excel Workbook (.xlsx) or as CSV, then import that.")
_LOCKED = ("is locked with a password, so it cannot be read. Open it in Excel, take the password off "
           "(File, Info, Protect Workbook, Encrypt with Password, and clear the box), save it, and "
           "import it again.")


class _Book:
    """The sheets of an .xls or .xlsx, read one at a time and only as far as asked."""

    def __init__(self, raw: bytes, filename: str):
        self.this = filename or "This file"
        if raw.startswith(XLS_MAGIC):
            try:
                import xlrd
            except ImportError as error:                  # pragma: no cover, it ships with the app
                raise RosterError("This build cannot read .xls files. Open the file in Excel and save "
                                  "it as CSV, then import that.") from error
            try:
                self._xls = xlrd.open_workbook(file_contents=raw)
            except Exception as error:
                if "encrypt" in str(error).lower():
                    raise RosterError(f"{self.this} {_LOCKED}") from error
                raise RosterError(_UNREADABLE) from error
            self.names = self._xls.sheet_names()
        else:
            try:
                from openpyxl import load_workbook
                self._xlsx = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
                self.names = list(self._xlsx.sheetnames)
            except Exception as error:
                raise RosterError(_UNREADABLE) from error

    def rows(self, name: str, limit: int | None = None) -> list[list[str]]:
        try:
            if hasattr(self, "_xls"):
                sheet = self._xls.sheet_by_name(name)
                count = sheet.nrows if limit is None else min(limit, sheet.nrows)
                return [[_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(count)]
            return [[_cell(v) for v in row]
                    for row in self._xlsx[name].iter_rows(values_only=True, max_row=limit)]
        except Exception as error:
            raise RosterError(_UNREADABLE) from error


def read_xls(raw: bytes) -> list[list[str]]:
    """The old binary Excel format, which is what a university portal hands out."""
    book = _Book(raw, "")
    return book.rows(book.names[0])


def sheet_for(names: list[str], level_number: int | None) -> str | None:
    """Which sheet of a workbook holds this level, by its name. None when no name says.

    A department keeps every year in one workbook, a sheet each: "200 Level", "300L", "400 Level ",
    "500 Level", "600 Level". Taking the first sheet would have imported 200 Level's students into
    whatever level the person was on (reported 2026-09-22). A sheet is matched on the number in its
    name and nothing else, because the rest of it is spelt differently on every tab.
    """
    if level_number is None:
        return None
    wanted = str(level_number)
    matches = [name for name in names
               if re.sub(r"[^0-9]", "", name) == wanted]
    return matches[0] if len(matches) == 1 else None


def _holds_a_list(table: list[list[str]]) -> bool:
    at = find_header(table)
    return at >= 0 and any(any(c.strip() for c in _trim(r or [])) for r in table[at + 1:])


def choose_sheet(book: _Book, level_number: int | None, asked: str | None = None) -> tuple[str, list[list[str]]]:
    """The sheet to read, and its rows.

    The one the person chose; else the one named for the level; else the only sheet; else the only
    sheet with a list on it, which is a workbook with a cover page in front of its list, refused
    until 2026-09-24. When several could be the list, the person is asked which.
    """
    if asked in book.names:
        return asked, book.rows(asked)
    named = sheet_for(book.names, level_number)
    if named:
        return named, book.rows(named)
    if len(book.names) == 1:
        return book.names[0], book.rows(book.names[0])
    lists = [name for name in book.names if _holds_a_list(book.rows(name, limit=HEADER_SEARCH_ROWS + 40))]
    if len(lists) == 1:
        return lists[0], book.rows(lists[0])
    listed = ", ".join(book.names)
    first = (lists or book.names)[0]
    table = book.rows(first)
    which = f"{level_number} Level" if level_number is not None else "the students"
    reason = (f"This file has {len(book.names)} sheets and none of them is named for {which}: {listed}. "
              f"Choose the sheet that holds {which}.") if level_number is not None else \
        f"This file has {len(book.names)} sheets: {listed}. Choose the sheet that holds the students."
    raise NeedsColumns(reason, table, guess_columns(table, find_header(table)), book.names, first)


def read_xlsx(raw: bytes, level_number: int | None = None) -> list[list[str]]:
    return choose_sheet(_Book(raw, ""), level_number)[1]


def not_a_roster(raw: bytes, filename: str = "") -> RosterError | None:
    """The refusal for a file that is some other kind of document, saying which kind.

    A PDF or a photo used to be "not readable text, save it as CSV", and a Word document "could not
    be read as a spreadsheet" (2026-09-24): true, and no help to somebody who picked the wrong file.
    """
    this = filename or "This file"

    def refuse(what: str, advice: str) -> RosterError:
        return RosterError(f"{this} is {what}, not a spreadsheet. {advice}")

    in_a_table = ("If the list is a table in it, copy the table into Excel, save it as an Excel "
                  "Workbook (.xlsx), then import that.")
    if raw.startswith(b"%PDF"):
        return refuse("a PDF", "Herewise reads the list from the spreadsheet itself (.xlsx, .xls or "
                               ".csv). If the department sent a PDF, ask them for the Excel file.")
    for magic, what in PICTURES:
        if raw.startswith(magic):
            return refuse(what, "A photo of a list cannot be read. Import the spreadsheet the list "
                                "came from.")
    if raw.startswith(b"MZ"):
        return refuse("a program", "Choose the roster spreadsheet instead.")
    if raw.startswith(b"{\\rtf"):
        return refuse("a Rich Text document", in_a_table)
    if raw.startswith(ZIP_MAGIC):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                inside = archive.namelist()
                kind = archive.read("mimetype") if "mimetype" in inside else b""
        except (zipfile.BadZipFile, KeyError, OSError):
            return RosterError(f"{this} is damaged and could not be opened. Save it again from "
                               "Excel, then import it.")
        if any(n.startswith("xl/") for n in inside):
            return None
        if any(n.startswith("word/") for n in inside):
            return refuse("a Word document", in_a_table)
        if any(n.startswith("ppt/") for n in inside):
            return refuse("a PowerPoint presentation", "Import the spreadsheet the list came from.")
        if b"opendocument.spreadsheet" in kind:
            return refuse("a LibreOffice spreadsheet (.ods)", "Open it in LibreOffice Calc and save it "
                                                              "as an Excel Workbook (.xlsx), then import that.")
        return refuse("a zip folder", "Open it, take the roster spreadsheet out, and import that.")
    if raw.startswith(XLS_MAGIC):
        if "EncryptionInfo".encode("utf-16-le") in raw:
            return RosterError(f"{this} {_LOCKED}")
        if "WordDocument".encode("utf-16-le") in raw:
            return refuse("a Word document", in_a_table)
        if "PowerPoint Document".encode("utf-16-le") in raw:
            return refuse("a PowerPoint presentation", "Import the spreadsheet the list came from.")
    return None


def _text_of(raw: bytes, filename: str) -> str:
    text = decode(raw)
    if "\x00" in text[:4096]:
        raise RosterError(f"{filename or 'This file'} is not a spreadsheet or a CSV file, so there is "
                          "nothing in it Herewise can read. Choose the roster itself (.xlsx, .xls or .csv).")
    return text


def parse_upload(raw: bytes, filename: str = "", level_number: int | None = None,
                 sheet: str | None = None, choice: Choice | None = None) -> list[dict]:
    """A roster as uploaded: a spreadsheet from the portal, a department's own, or a CSV.

    The format is decided by what is in the file, not by what it is called. Portals hand out files
    named .xls that are really a web page, and a person renaming a file does not change it.

    ``level_number`` is the level being imported into, used to choose a sheet when the workbook
    keeps one per year. ``sheet`` and ``choice`` are what the person picked when asked.
    """
    if len(raw) > MAX_UPLOAD_BYTES:
        size = len(raw) / (1024 * 1024)
        raise RosterError(f"{filename or 'This file'} is {size:.1f} MB. A roster is a few kilobytes (the "
                          "department's longest list is 16 KB), so this is probably not the list. The "
                          "most Herewise reads is 5 MB.")
    if not raw.strip():
        raise RosterError("The file is empty.")
    refusal = not_a_roster(raw, filename)
    if refusal:
        raise refusal
    sheets, chosen = [], None
    if raw.startswith(XLS_MAGIC) or raw.startswith(ZIP_MAGIC):
        book = _Book(raw, filename)
        sheets = book.names if len(book.names) > 1 else []
        chosen, table = choose_sheet(book, level_number, sheet)
    else:
        head = raw[:1024].lstrip().lower()
        if head.startswith(b"<") and (b"<html" in head or b"<?xml" in head or b"<table" in head):
            raise RosterError(
                f"{filename or 'This file'} is named like a spreadsheet but is really a web page the "
                "portal saved. Open it in Excel and save it as CSV or as an Excel Workbook, then "
                "import that.")
        table = csv_table(_text_of(raw, filename))
    if choice is not None:
        return rows_by_choice(table, choice)
    try:
        rows, at, columns = _read(table)
    except NeedsColumns as asking:
        asking.sheets, asking.sheet = sheets, chosen
        raise
    trouble = column_trouble(rows)
    if trouble:
        guess = {"matric_col": None, "name_col": None, "other_col": None, "level_col": None,
                 "first_row": at + 2}
        header = [column_name(h) for h in _trim(table[at] or [])]
        if "name" in header and "matric_no" in header:
            # The two the wrong way round is the likeliest reading, so it is the one offered.
            swapped = trouble.startswith("The matric number and name columns")
            guess["matric_col"] = header.index("name") if swapped else header.index("matric_no")
            guess["name_col"] = header.index("matric_no") if swapped else header.index("name")
        else:
            guess = guess_columns(table, at)
        if "level" in header:
            guess["level_col"] = header.index("level")
        raise NeedsColumns(f"{trouble} Say which column is which below.", table, guess, sheets, chosen)
    return rows


def as_csv(rows: list[dict]) -> str:
    """The parsed rows as CSV, for the record kept with the import.

    A spreadsheet cannot be stored in a text column, and the commit re-reads this text rather than
    the file, so what is committed is exactly what the preview showed. It also means the record of
    an import is readable a year later without the portal's own file.
    """
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\n")
    # The surname and other names are kept apart when the file kept them apart, so a surname of two
    # words is not split again by guesswork when the import is committed.
    writer.writerow(["matric_no", "name", "level", "surname", "other_names"])
    for row in rows:
        split = bool(row.get("surname"))
        writer.writerow([row.get("matric_no", ""), row.get("name", ""), row.get("level", ""),
                         row.get("surname", "") if split else "", row.get("other_names", "") if split else ""])
    return out.getvalue()



# The template's heading row. Everything above it is instructions, which the reader never reaches
# because it starts from the headings (find_header), so the example cannot be imported by mistake.
TEMPLATE_HEADINGS = ("MATRIC NO", "SURNAME", "OTHER NAMES")
TEMPLATE_HEADING_ROW = 5


def template(level_name: str) -> bytes:
    """A blank roster for a list typed in by hand, as an Excel workbook.

    The sheet is named for the level, so a template filled in for 200 Level and imported into it is
    read without a question. The surname has a column of its own because a typed list is where a
    two-word surname would otherwise be split in the wrong place.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = level_name[:31]
    sheet["A1"] = f"{level_name} roster"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (f"One student per row, under the headings on row {TEMPLATE_HEADING_ROW}. Keep the "
                   "headings as they are.")
    sheet["A3"] = ("For example: 25/PHARX01/001 under MATRIC NO, ADEYEMI under SURNAME, Tolu Grace "
                   "under OTHER NAMES. This row is not read.")
    sheet["A4"] = "A student with no matric number yet: type DE under MATRIC NO."
    for column, heading in enumerate(TEMPLATE_HEADINGS, start=1):
        cell = sheet.cell(row=TEMPLATE_HEADING_ROW, column=column, value=heading)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7E2DA")
    for letter, width in zip("ABC", (20, 24, 36)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = f"A{TEMPLATE_HEADING_ROW + 1}"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def next_device_id(conn: sqlite3.Connection, level: sqlite3.Row) -> int:
    """The next ID in a level's block. IDs are never reused, because old scans refer to them.

    Scans count as well as people: a scan from an ID nobody here has (someone added at the keypad, a
    user left from an earlier setup) would otherwise be handed to the next new student as theirs. So
    do the IDs of a deleted roster (table retired_ids), whose scans and people are gone.
    """
    start, end = level["id_block_start"], level["id_block_end"]
    in_people = conn.execute(
        "SELECT COALESCE(MAX(device_user_id), 0) FROM people WHERE device_user_id BETWEEN ? AND ?",
        (start, end),
    ).fetchone()[0]
    # punches.device_user_id is text, straight from the terminal; only whole numbers can clash.
    in_scans = conn.execute(
        "SELECT COALESCE(MAX(CAST(device_user_id AS INTEGER)), 0) FROM punches"
        " WHERE device_user_id <> '' AND device_user_id NOT GLOB '*[^0-9]*'"
        " AND CAST(device_user_id AS INTEGER) BETWEEN ? AND ?",
        (start, end),
    ).fetchone()[0]
    # A deleted roster's IDs are retired, not freed: the terminal may still log a scan under one.
    retired = levels.highest_retired(conn, start, end)
    return max(in_people + 1, in_scans + 1, retired + 1, start)


def build_plan(conn: sqlite3.Connection, level_id: int, academic_session_id: int, device_id: int,
               rows: list[dict]) -> dict:
    level = conn.execute("SELECT * FROM levels WHERE id = ?", (level_id,)).fetchone()
    if level is None:
        raise RosterError("Unknown level.")
    device = conn.execute("SELECT * FROM devices WHERE id = ?", (device_id,)).fetchone()
    if device is None:
        raise RosterError("Unknown device.")

    by_matric = {r["matric_no"]: r for r in conn.execute("SELECT * FROM people WHERE matric_no IS NOT NULL")}
    # Direct-entry students have no matric number yet (reported 2026-09-23: "those are direct entry
    # students that don't have matric numbers yet"). They are imported all the same, and recognised
    # again by their name within the level, so importing the same list twice does not make a second
    # copy, and a later list that carries their number fills it in rather than adding a stranger.
    without_matric = {
        names.fold(f"{r['surname']} {r['other_names']}".strip()): r for r in conn.execute(
            "SELECT p.* FROM people p JOIN level_memberships m ON m.person_id = p.id"
            " WHERE p.matric_no IS NULL AND m.level_id = ? AND m.academic_session_id = ?",
            (level_id, academic_session_id))}
    seen_names: dict[str, int] = {}
    memberships = {
        r["person_id"]: r for r in conn.execute(
            "SELECT m.person_id, l.name AS level_name, l.id AS level_id FROM level_memberships m"
            " JOIN levels l ON l.id = m.level_id WHERE m.academic_session_id = ?", (academic_session_id,))
    }
    next_id = next_device_id(conn, level)

    planned: list[PlannedStudent] = []
    seen: dict[str, int] = {}
    errors: list[str] = []

    for row in rows:
        line = row["_line"]
        matric = names.matric_or_none(row.get("matric_no")) or ""
        raw_name = row.get("name", "")
        if row.get("surname"):
            # The file kept the surname in a column of its own, so it is taken as it is.
            parsed = names.ParsedName(names.clean(row["surname"]), names.clean(row.get("other_names", "")))
        else:
            parsed = names.split_name(raw_name)
        dname = names.device_name(parsed.surname, parsed.other_names)
        full = " ".join(p for p in (parsed.surname, parsed.other_names) if p)
        item = PlannedStudent(line, matric, raw_name, parsed.surname, parsed.other_names, dname,
                              shortened=(dname != full), action="new", issues=list(parsed.warnings))
        if row.get("_joined"):
            item.issues.append("name contained an unquoted comma, columns rejoined")

        key = names.fold(full)
        if not parsed.surname:
            # One reason, not two: "name is empty" beside "missing name" read as one run-on phrase.
            item.action = "skip"
            item.issues = ["the name is empty"]
        elif any(ch.isdigit() for ch in full):
            # A total under the list, a serial number in the wrong column, or a matric number: never
            # a person's name, and it used to be imported as a student without a matric number.
            item.action = "skip"
            item.issues.append("the name has a number in it, so it is not a student's name")
        elif matric and not any(ch.isdigit() for ch in matric):
            item.action = "skip"
            item.issues.append("the matric number has no number in it (type DE, or leave it empty, "
                               "for a student who has none yet)")
        elif not matric:
            if key in seen_names:
                item.action = "skip"
                item.issues.append(f"the same name with no matric number on line {seen_names[key]}")
            else:
                seen_names[key] = line
                item.issues.append("direct entry (DE): imported without a matric number, which can be "
                                   "added later")
        elif matric in seen:
            item.action = "skip"
            item.issues.append(f"duplicate matric number, first seen on line {seen[matric]}")
        else:
            seen[matric] = line

        level_value = row.get("level", "").strip()
        if item.action != "skip" and level_value and level_value != str(level["level_number"]):
            item.action = "skip"
            item.issues.append(f"row says level {level_value}, importing into {level['level_number']}")

        if item.action != "skip":
            existing = by_matric.get(matric) if matric else None
            if existing is None and key in without_matric:
                # The same student, known so far without a number: this row fills it in, or is them.
                existing = without_matric[key]
                if matric:
                    item.issues.append("matric number added to a student who had none")
            if existing is not None:
                item.person_id = existing["id"]
                item.device_user_id = existing["device_user_id"]
                membership = memberships.get(existing["id"])
                changed_name = ((existing["surname"], existing["other_names"]) != (parsed.surname, parsed.other_names)
                                or (bool(matric) and existing["matric_no"] != matric))
                if membership and membership["level_id"] != level_id:
                    item.issues.append(f"moves from {membership['level_name']} this session")
                item.action = "update" if (changed_name or not membership or membership["level_id"] != level_id) \
                    else "unchanged"
                if existing["device_id"] != device_id:
                    item.issues.append("already assigned to another device, keeps that device")
            else:
                if next_id > level["id_block_end"]:
                    item.action = "skip"
                    item.issues.append(f"no IDs left in block {level['id_block_start']}-{level['id_block_end']}")
                else:
                    item.device_user_id = next_id
                    next_id += 1
        planned.append(item)

    new = [p for p in planned if p.action == "new"]
    assigned = conn.execute("SELECT COUNT(*) FROM people WHERE device_id = ? AND active = 1", (device_id,)).fetchone()[0]
    capacity = device["user_capacity"] or DEFAULT_USER_CAPACITY
    after = assigned + len(new)
    if after > capacity:
        errors.append(f"This import would put {after} people on {device['name']}, which holds {capacity}. "
                      "Assign this level to another device.")

    return {
        "level": {"id": level["id"], "name": level["name"]},
        "warnings": _warnings(planned, level),
        "rows": len(planned),
        "new": len(new),
        "updated": sum(1 for p in planned if p.action == "update"),
        "unchanged": sum(1 for p in planned if p.action == "unchanged"),
        "skipped": sum(1 for p in planned if p.action == "skip"),
        "id_range": [new[0].device_user_id, new[-1].device_user_id] if new else None,
        "shortened": [{"line": p.line, "before": " ".join(x for x in (p.surname, p.other_names) if x),
                       "after": p.device_name} for p in planned if p.shortened],
        "issues": [{"line": p.line, "matric_no": p.matric_no, "action": p.action, "issues": p.issues}
                   for p in planned if p.issues],
        "capacity": {"device": device["name"], "assigned_now": assigned, "after_import": after,
                     "capacity": capacity},
        "errors": errors,
        "can_commit": not errors and any(p.action in ("new", "update") for p in planned),
        "students": [asdict(p) for p in planned],
    }


def _warnings(planned: list[PlannedStudent], level: sqlite3.Row) -> list[dict]:
    """What has to be ticked before this import goes ahead, each with the sentence and the box.

    Neither stops an import. Each is a file that can be imported and is probably not the one meant:
    300 Level's list, imported into 200 Level with no level column, offered to move all six of them
    with the news in each row's own line; and a file with five rows of six refused still offered
    "Import 1 student" (2026-09-24). The commit refuses until the same warnings are confirmed.
    """
    said = []
    importable = [p for p in planned if p.action != "skip"]
    moving = [p for p in importable if any(i.startswith("moves from ") for i in p.issues)]
    if len(moving) >= 3 and len(moving) * 2 > len(importable):
        homes = Counter(i[len("moves from "):].removesuffix(" this session")
                        for p in moving for i in p.issues if i.startswith("moves from "))
        home = homes.most_common(1)[0][0]
        said.append({
            "code": "moves",
            "text": (f"{len(moving)} of the {len(planned)} students in this file are on {home} already, "
                     f"and importing moves them to {level['name']}. If this is {home}'s list, import "
                     "it there instead."),
            "confirm": f"Yes, move them to {level['name']}",
        })
    refused = len(planned) - len(importable)
    if importable and refused * 2 > len(planned):
        said.append({
            "code": "mostly_refused",
            "text": (f"{refused} of the {len(planned)} rows in this file are refused, so it may be the "
                     "wrong file, or its columns may not hold what they seem to. Read the reasons "
                     "before importing the rest."),
            "confirm": f"Import the {len(importable)} that can be imported anyway",
        })
    return said


def _tracking_from(conn: sqlite3.Connection) -> str:
    """When Herewise began recording which import created a student. Read straight from settings,
    because db.get_setting insists on a default and there is no sensible default for a date."""
    row = conn.execute("SELECT value FROM settings WHERE key = 'roster_tracking_from'").fetchone()
    return row["value"] if row else ""


def imports_for(conn: sqlite3.Connection, level_id: int) -> list[dict]:
    """Every import for a level, newest first, with what it did and whether it can be undone."""
    rows = conn.execute(
        "SELECT i.*, a.display_name AS by_whom,"
        " (SELECT COUNT(*) FROM people p WHERE p.roster_import_id = i.id) AS people_created"
        " FROM roster_imports i LEFT JOIN accounts a ON a.id = i.account_id"
        " WHERE i.level_id = ? AND i.status = 'committed' ORDER BY i.id DESC", (level_id,)).fetchall()
    tracked_from = _tracking_from(conn)
    out = []
    for row in rows:
        held = _attached(conn, row["id"])
        when = row["committed_at"] or row["created_at"] or ""
        out.append({**{k: row[k] for k in row.keys() if k != "csv_text"},
                    "can_undo": row["people_created"] > 0 and not held,
                    "before_tracking": bool(row["people_created"] == 0 and tracked_from
                                            and when < tracked_from),
                    "holding": held})
    return out


def _attached(conn: sqlite3.Connection, import_id: int) -> dict[str, int]:
    """What has happened to the students an import created, by kind.

    An import is only undone while its students are untouched: a scan, a correction, an excused
    absence, a fingerprint or a keypad password all mean somebody has been at the terminal or the
    register since, and a device ID that has been scanned against can never be handed back.
    """
    counts = {
        "scans": "SELECT COUNT(*) FROM punches WHERE device_user_id IN"
                 " (SELECT CAST(device_user_id AS TEXT) FROM people WHERE roster_import_id = ?)",
        "corrections": "SELECT COUNT(*) FROM attendance_marks WHERE person_id IN"
                       " (SELECT id FROM people WHERE roster_import_id = ?)",
        "excused absences": "SELECT COUNT(*) FROM excusals WHERE person_id IN"
                            " (SELECT id FROM people WHERE roster_import_id = ?)",
        "fingerprints": "SELECT COUNT(*) FROM fingers WHERE person_id IN"
                        " (SELECT id FROM people WHERE roster_import_id = ?)",
        "on the terminal": "SELECT COUNT(*) FROM people WHERE roster_import_id = ? AND on_device = 1",
        "keypad passwords": "SELECT COUNT(*) FROM people WHERE roster_import_id = ?"
                            " AND device_password IS NOT NULL",
    }
    held = {}
    for what, sql in counts.items():
        found = conn.execute(sql, (import_id,)).fetchone()[0]
        if found:
            held[what] = found
    return held


def undo(conn: sqlite3.Connection, import_id: int, account_id: int | None) -> dict:
    """Take back an import: remove the students it created, if nothing has happened to them.

    Asked for on a first run, after a roster went into the wrong level: there was no way back short
    of editing the database. Students the import only updated are left alone, because their earlier
    import is what put them there.
    """
    record = conn.execute("SELECT * FROM roster_imports WHERE id = ?", (import_id,)).fetchone()
    if record is None:
        raise RosterError("Unknown import.")
    if record["status"] != "committed":
        raise RosterError(f"That import is {record['status']}, so there is nothing to take back.")
    created = conn.execute("SELECT id, matric_no FROM people WHERE roster_import_id = ?",
                           (import_id,)).fetchall()
    if not created:
        tracked_from = _tracking_from(conn)
        when = record["committed_at"] or record["created_at"] or ""
        if tracked_from and when < tracked_from:
            raise RosterError(
                "This import was made before Herewise began recording which students an import "
                "created, so the students it created cannot be told from any others. Remove them "
                "one at a time from their own records instead.")
        raise RosterError("This import created nobody, so there is nothing to take back. It only "
                          "updated students who were already here.")
    held = _attached(conn, import_id)
    if held:
        what = ", ".join(f"{count} {name}" for name, count in sorted(held.items()))
        raise RosterError(
            f"These {len(created)} students have {what}, so the import cannot be taken back. "
            "A device ID that has been scanned against is never given to anybody else, and "
            "attendance already recorded is never deleted.")
    with db.transaction(conn):
        ids = [r["id"] for r in created]
        marks = ",".join("?" * len(ids))
        conn.execute(f"DELETE FROM level_memberships WHERE person_id IN ({marks})", ids)
        conn.execute(f"DELETE FROM course_enrolments WHERE person_id IN ({marks})", ids)
        conn.execute(f"DELETE FROM course_group_members WHERE person_id IN ({marks})", ids)
        conn.execute(f"DELETE FROM people WHERE id IN ({marks})", ids)
        conn.execute("UPDATE roster_imports SET status = 'discarded' WHERE id = ?", (import_id,))
        db.audit(conn, account_id, "roster.undo", "roster_imports", import_id,
                 f"{len(ids)} students taken back out, from {record['file_name'] or 'an import'}")
    return {"removed": len(created), "import_id": import_id}


def edit_student(conn: sqlite3.Connection, person_id: int, surname: str, other_names: str,
                 matric_no: str | None, account_id: int | None) -> dict:
    """Correct a student's name or matric number, or give them the number they did not have.

    Asked for on 2026-09-23: direct-entry students have no matric number when the list is made, so the
    names and matric numbers on the sheet have to be editable, "so that they can be added in the
    future". The device ID never changes, so scans, fingerprints and corrections stay with the student.
    The name the terminal shows is worked out again; ``device_name_changed`` tells the caller to send it.
    """
    person = conn.execute("SELECT * FROM people WHERE id = ? AND role = 'student'", (person_id,)).fetchone()
    if person is None:
        raise RosterError("Unknown student.")
    surname = " ".join((surname or "").split())
    other_names = " ".join((other_names or "").split())
    if not surname:
        raise RosterError("A student needs a surname.")
    matric = names.matric_or_none(matric_no)
    if matric:
        taken = conn.execute("SELECT surname, other_names FROM people WHERE matric_no = ? AND id != ?",
                             (matric, person_id)).fetchone()
        if taken is not None:
            raise RosterError(f"{matric} already belongs to {taken['surname']} {taken['other_names']}".strip() + ".")
    dname = names.device_name(surname, other_names)
    changes = []
    if (person["surname"], person["other_names"]) != (surname, other_names):
        changes.append(f"name {person['surname']} {person['other_names']} to {surname} {other_names}".replace("  ", " "))
    if person["matric_no"] != matric:
        changes.append(f"matric number {person['matric_no'] or 'none'} to {matric or 'none'}")
    if not changes:
        return {"person_id": person_id, "changed": False, "device_name_changed": False}
    with db.transaction(conn):
        conn.execute(
            "UPDATE people SET surname = ?, other_names = ?, matric_no = ?, device_name = ?,"
            " on_device = CASE WHEN device_name = ? THEN on_device ELSE 0 END WHERE id = ?",
            (surname, other_names, matric, dname, dname, person_id))
        db.audit(conn, account_id, "person.edit", "people", person_id, "; ".join(changes))
    return {"person_id": person_id, "changed": True, "device_name_changed": dname != person["device_name"],
            "was_on_device": bool(person["on_device"]), "device_id": person["device_id"]}


def commit(conn: sqlite3.Connection, import_id: int, account_id: int | None,
           confirmed: Iterable[str] = ()) -> dict:
    with db.transaction(conn):
        record = conn.execute("SELECT * FROM roster_imports WHERE id = ?", (import_id,)).fetchone()
        if record is None:
            raise RosterError("Unknown import.")
        if record["status"] != "pending":
            raise RosterError(f"This import is already {record['status']}.")
        plan = build_plan(conn, record["level_id"], record["academic_session_id"], record["device_id"],
                          parse_csv(record["csv_text"]))
        if not plan["can_commit"]:
            raise RosterError("; ".join(plan["errors"]) or "Nothing to import.")
        unconfirmed = [w for w in plan["warnings"] if w["code"] not in set(confirmed)]
        if unconfirmed:
            raise RosterError(f"{unconfirmed[0]['text']} Tick the box under it to go ahead.")

        stamp = db.now()
        for s in plan["students"]:
            if s["action"] == "new":
                cur = conn.execute(
                    "INSERT INTO people(role, device_user_id, device_id, matric_no, surname, other_names,"
                    " source_name, device_name, created_at, roster_import_id)"
                    " VALUES ('student', ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    # NULL, never "": matric_no is unique, and two empty strings would collide.
                    (s["device_user_id"], record["device_id"], s["matric_no"] or None, s["surname"], s["other_names"],
                     s["source_name"], s["device_name"], stamp, import_id),
                )
                person_id = cur.lastrowid
            elif s["action"] == "update":
                person_id = s["person_id"]
                # on_device drops to 0 only if the device name changed, so the next push updates it.
                conn.execute(
                    "UPDATE people SET surname = ?, other_names = ?, source_name = ?, device_name = ?,"
                    " on_device = CASE WHEN device_name = ? THEN on_device ELSE 0 END, active = 1,"
                    " matric_no = coalesce(?, matric_no) WHERE id = ?",
                    (s["surname"], s["other_names"], s["source_name"], s["device_name"], s["device_name"],
                     s["matric_no"] or None, person_id),
                )
            else:
                continue
            conn.execute(
                "INSERT INTO level_memberships(person_id, academic_session_id, level_id) VALUES (?, ?, ?)"
                " ON CONFLICT(person_id, academic_session_id) DO UPDATE SET level_id = excluded.level_id",
                (person_id, record["academic_session_id"], record["level_id"]),
            )
        conn.execute("UPDATE roster_imports SET status = 'committed', committed_at = ?, account_id = ? WHERE id = ?",
                     (stamp, account_id, import_id))
        db.audit(conn, account_id, "roster.commit", "roster_imports", import_id,
                 f"{plan['new']} new, {plan['updated']} updated, {plan['skipped']} skipped into {plan['level']['name']}")
        # Every imported list is on its level's courses straight away. A level with no course yet gets
        # none: the lecturer makes it, and it takes the whole level then.
        course_ids = courses.on_roster_commit(conn, record["level_id"], record["academic_session_id"], account_id)
        template = conn.execute("SELECT id FROM register_templates WHERE course_id = ?",
                                (course_ids[0],)).fetchone() if course_ids else None
        template_id = template["id"] if template else None
    summary = {k: v for k, v in plan.items() if k != "students"}
    summary["course_ids"] = course_ids
    summary["template_id"] = template_id
    return summary
